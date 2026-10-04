"""The drone patrol report: what flew, where, what it saw and what was done.

ONE LOADER, THREE VIEWS. `load_report` reads a flight into a plain dictionary;
the PDF, the workbook and the JSON endpoint are all that dictionary rendered. A
number in the PDF cannot disagree with the same number in the spreadsheet,
because neither computes it.

EVERY VALUE IS THE FLIGHT'S OWN. Mission, route, drone and profile names, the
waypoints and the zones come from the session and its frozen configuration, not
from the live tables — a mission renamed or a route redrawn next month does not
rewrite last month's report.

THE PICTURE IS THE ONE TAKEN THEN. An event's snapshot is the file recorded with
it: the gateway's upload, or the AI worker's own evidence for the detection.
Nothing is fetched from a camera at report time. A file still at the site, or
no longer in storage, is said to be so — a report that silently drops a missing
picture makes a lost one look like one that was never taken.

AI CONFIDENCE AND RISK STAY TWO COLUMNS, as everywhere else in this module.

TIMES ARE THE ORGANISATION'S OWN CLOCK, labelled with the zone. A report read by
a site manager says 22:14, not 14:14 UTC.

The libraries are the ones the platform's other reports use (reportlab,
openpyxl), imported lazily and guarded the same way.
"""
from __future__ import annotations

import asyncio
import hashlib
import io
import logging
import math
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from typing import Callable
from xml.sax.saxutils import escape
from zoneinfo import ZoneInfo

from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.services.drone_ai_pipeline import tenant_tz

logger = logging.getLogger(__name__)

RISK_LEVELS = ("CRITICAL", "HIGH", "MEDIUM", "LOW", "INFO")
#: "Suspicious" in the summary: assessed MEDIUM or above and not dismissed.
SUSPICIOUS = ("MEDIUM", "HIGH", "CRITICAL")
TERMINAL = ("COMPLETED", "FAILED", "ABORTED", "CANCELLED", "BLOCKED", "MISSED")
#: The flown track is drawn, not listed: this many points show its shape.
MAX_TRACK_POINTS = 400
MAX_IMAGE_WIDTH_MM = 120
MAX_IMAGE_HEIGHT_MM = 75
RISK_HEX = {"INFO": "#8EA0B8", "LOW": "#00A878", "MEDIUM": "#E09600", "HIGH": "#E8642C", "CRITICAL": "#D9263F"}
#: The same levels as a pale cell background, for tables on paper and in Excel.
RISK_TINT = {"INFO": "E8ECF2", "LOW": "D9F2E9", "MEDIUM": "FBEBC8", "HIGH": "FBDDCF", "CRITICAL": "F9D0D6"}

#: What each AI module is called on paper. "Lpr" is not a word a client reads.
MODULE_NAMES = {"lpr": "Licence plate", "face": "Face", "intrusion": "Intrusion", "ppe": "PPE", "crowd": "Crowd",
                "fire_smoke": "Fire or smoke", "weapon": "Weapon", "behavior": "Behaviour",
                "tampering": "Camera tampering", "abandoned": "Abandoned object", "fall": "Fall"}

FileReader = Callable[[str], bytes | None]


def _check_reportlab():
    try:
        import reportlab  # noqa: F401
    except ImportError:  # pragma: no cover - environment guard
        raise RuntimeError("reportlab is not installed — rebuild the container with reportlab>=4.2")


def _check_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ImportError:  # pragma: no cover - environment guard
        raise RuntimeError("openpyxl is not installed - rebuild the container with openpyxl>=3.1")


def _f(v) -> float | None:
    return float(v) if v is not None else None


def _pretty(s: str | None) -> str:
    t = (s or "").replace("_", " ").lower()
    return t[:1].upper() + t[1:]


# ═════════════════════════════════════════════════════════════════════════════
# Loading
# ═════════════════════════════════════════════════════════════════════════════

def officer_action(e: dict) -> str:
    """What a person did about an event, in a sentence."""
    status = e["status"]
    if status == "FALSE_POSITIVE":
        who = f" by {e['resolved_by_name']}" if e.get("resolved_by_name") else ""
        return f"Marked a false positive{who}: {e.get('false_positive_reason') or 'no reason recorded'}"
    if status == "RESOLVED":
        return f"Resolved by {e['resolved_by_name']}" if e.get("resolved_by_name") else "Resolved"
    if status == "ESCALATED":
        return "Escalated"
    if status == "INVESTIGATING":
        return "Under investigation"
    if status == "ACKNOWLEDGED":
        return f"Acknowledged by {e['acknowledged_by_name']}" if e.get("acknowledged_by_name") else "Acknowledged"
    return "No action recorded"


def _snapshot(media: list[dict], evidence: list[dict]) -> dict:
    """The event's picture and where it is: the gateway's upload first, then the
    AI worker's own evidence for the detection."""
    shots = [m for m in media if m["media_kind"] == "SNAPSHOT"]
    for m in shots:
        if m["storage_location"] in ("central", "both"):
            return {"state": "available", "source": "drone", "path": m["storage_path"],
                    "checksum_sha256": m["checksum_sha256"], "captured_at": m["captured_at"]}
    if evidence:
        ev = evidence[0]
        return {"state": "available", "source": "ai_worker", "path": ev["storage_path"],
                "checksum_sha256": ev["checksum_sha256"], "captured_at": ev["captured_at"]}
    if shots:
        return {"state": "held_at_site", "source": "drone", "path": None,
                "checksum_sha256": shots[0]["checksum_sha256"], "captured_at": shots[0]["captured_at"],
                "sync_state": shots[0]["sync_state"]}
    return {"state": "none", "source": None, "path": None, "checksum_sha256": None, "captured_at": None}


async def load_report(db: AsyncSession, session_id) -> dict:
    """Everything a flight's report shows. Raises ValueError if there is no such
    flight in the tenant the connection is scoped to."""
    sid = str(session_id)
    s = (await db.execute(text("""
        SELECT ps.*, si.name AS site_name, tu.full_name AS triggered_by_name, au.full_name AS aborted_by_name,
               d.code AS drone_code
          FROM drone_patrol_sessions ps
          LEFT JOIN sites si ON si.id = ps.site_id
          LEFT JOIN users tu ON tu.id = ps.triggered_by_user_id
          LEFT JOIN users au ON au.id = ps.aborted_by_user_id
          LEFT JOIN drones d ON d.id = ps.drone_id
         WHERE ps.id = CAST(:id AS uuid)
    """), {"id": sid})).mappings().first()
    if s is None:
        raise ValueError("Patrol session not found")
    tz = await tenant_tz(db)
    snap = s["config_snapshot"] or {}

    waypoints = (await db.execute(text("""
        SELECT sequence, name, latitude, longitude, altitude_m, hover_seconds, observe_seconds,
               snapshot_required, status, reached_at, departed_at
          FROM drone_session_waypoints WHERE session_id = CAST(:id AS uuid) ORDER BY sequence
    """), {"id": sid})).mappings().all()

    total = (await db.execute(text(
        "SELECT count(*) FROM drone_telemetry WHERE session_id = CAST(:id AS uuid) AND latitude IS NOT NULL"),
        {"id": sid})).scalar() or 0
    step = max(1, -(-total // MAX_TRACK_POINTS))
    track = (await db.execute(text("""
        SELECT latitude, longitude, battery_pct, recorded_at
          FROM (SELECT t.*, row_number() OVER (ORDER BY recorded_at) AS rn
                  FROM drone_telemetry t
                 WHERE t.session_id = CAST(:id AS uuid) AND t.latitude IS NOT NULL) x
         WHERE (rn - 1) % :step = 0 OR rn = :total
         ORDER BY recorded_at
    """), {"id": sid, "step": step, "total": total})).mappings().all()

    events = (await db.execute(text("""
        SELECT e.*, au.full_name AS acknowledged_by_name, ru.full_name AS resolved_by_name,
               (SELECT count(*) FROM drone_observations o WHERE o.event_id = e.id) AS observation_count
          FROM drone_events e
          LEFT JOIN users au ON au.id = e.acknowledged_by_user_id
          LEFT JOIN users ru ON ru.id = e.resolved_by_user_id
         WHERE e.session_id = CAST(:id AS uuid)
         ORDER BY e.detected_at
    """), {"id": sid})).mappings().all()
    event_ids = [e["id"] for e in events]

    media: dict = {}
    cameras: dict = {}
    evidence: dict = {}
    incidents: dict = {}
    if event_ids:
        for m in (await db.execute(text("""
            SELECT id, event_id, media_kind, storage_path, storage_location, sync_state, checksum_sha256,
                   size_bytes, duration_seconds, captured_at
              FROM drone_event_media WHERE event_id = ANY(:ids) ORDER BY captured_at
        """), {"ids": event_ids})).mappings().all():
            media.setdefault(m["event_id"], []).append(dict(m))
        for c in (await db.execute(text("""
            SELECT event_id, camera_name, distance_m, correlation_method, corroborates,
                   related_detection_count, related_module_type, recording_id
              FROM drone_event_cameras WHERE event_id = ANY(:ids)
             ORDER BY rank NULLS LAST, distance_m NULLS LAST
        """), {"ids": event_ids})).mappings().all():
            cameras.setdefault(c["event_id"], []).append(dict(c))
        # The AI worker's own snapshot of a detection, where the platform keeps
        # one. Read only: `evidence` is the platform's table.
        for ev in (await db.execute(text("""
            SELECT o.event_id, ev.storage_path, ev.checksum_sha256, ev.captured_at
              FROM drone_observations o
              JOIN evidence ev ON ev.detection_id = o.detection_id AND ev.media_type = 'image'
             WHERE o.event_id = ANY(:ids) AND o.detection_id IS NOT NULL
             ORDER BY ev.captured_at
        """), {"ids": event_ids})).mappings().all():
            evidence.setdefault(ev["event_id"], []).append(dict(ev))
        incident_ids = sorted({e["incident_id"] for e in events if e["incident_id"]}, key=str)
        if incident_ids:
            for i in (await db.execute(text("""
                SELECT i.id, i.title, i.severity, i.status, i.created_at, i.resolved_at, i.dispatched_at,
                       i.guard_arrived_at, i.message_params->>'incident_ref' AS incident_ref,
                       g.full_name AS dispatched_guard_name
                  FROM incidents i LEFT JOIN users g ON g.id = i.dispatched_guard_id
                 WHERE i.id = ANY(:ids)
            """), {"ids": incident_ids})).mappings().all():
                incidents[i["id"]] = dict(i)

    started = s["launched_at"] or s["started_at"]
    flew = s["launched_at"] is not None
    duration = (s["ended_at"] - started).total_seconds() if (started and s["ended_at"]) else None
    route = snap.get("route") or {}
    base = ({"latitude": route["base_latitude"], "longitude": route["base_longitude"]}
            if route.get("base_latitude") is not None and route.get("base_longitude") is not None else None)

    out_events, by_risk, by_module = [], {level: 0 for level in RISK_LEVELS}, {}
    for n, e in enumerate(events, start=1):
        e = dict(e)
        by_risk[e["risk_level"]] = by_risk.get(e["risk_level"], 0) + 1
        by_module[e["module_type"]] = by_module.get(e["module_type"], 0) + 1
        ms = media.get(e["id"], [])
        inc = incidents.get(e["incident_id"])
        out_events.append({
            "number": n, "id": str(e["id"]), "detected_at": e["detected_at"],
            "last_detected_at": e.get("last_detected_at"),
            "module_type": e["module_type"], "label": e.get("label"),
            "latitude": e["drone_latitude"], "longitude": e["drone_longitude"],
            "location_method": e["location_method"], "zone_name": e["zone_name"], "zone_type": e["zone_type"],
            "waypoint_sequence": e["waypoint_sequence"],
            "ai_confidence": _f(e["ai_confidence"]), "risk_score": e["risk_score"], "risk_level": e["risk_level"],
            "risk_factors": e["risk_factors"] or [],
            "detections": e["observation_count"] or e.get("detection_count") or 0,
            "observed_seconds": _f(e["observed_seconds"]),
            "verification_state": e["verification_state"], "status": e["status"],
            "officer_action": officer_action(e),
            "snapshot": _snapshot(ms, evidence.get(e["id"], [])),
            "videos": [{"media_id": str(m["id"]), "kind": m["media_kind"], "location": m["storage_location"],
                        "sync_state": m["sync_state"], "checksum_sha256": m["checksum_sha256"],
                        "size_bytes": m["size_bytes"], "duration_seconds": _f(m["duration_seconds"]),
                        "captured_at": m["captured_at"]}
                       for m in ms if m["media_kind"] != "SNAPSHOT"],
            "cctv": [{"camera_name": c["camera_name"], "distance_m": _f(c["distance_m"]),
                      "method": c["correlation_method"], "corroborates": bool(c["corroborates"]),
                      "detections": c["related_detection_count"] or 0, "module_type": c["related_module_type"],
                      "has_recording": c["recording_id"] is not None}
                     for c in cameras.get(e["id"], [])],
            "incident": ({"id": str(inc["id"]), "ref": inc["incident_ref"], "title": inc["title"],
                          "severity": inc["severity"], "status": inc["status"],
                          "dispatched_guard_name": inc["dispatched_guard_name"],
                          "dispatched_at": inc["dispatched_at"], "guard_arrived_at": inc["guard_arrived_at"],
                          "resolved_at": inc["resolved_at"]} if inc else None),
        })

    return {
        "session_id": sid, "tenant_id": str(s["tenant_id"]), "site_id": str(s["site_id"]) if s["site_id"] else None,
        "generated_at": datetime.now(timezone.utc), "timezone": tz.key,
        "mission": {
            "session_number": s["session_number"], "site_name": s["site_name"], "mission_name": s["mission_name"],
            "drone_name": s["drone_name"], "drone_code": s["drone_code"], "route_name": s["route_name"],
            "profile_name": s["profile_name"], "triggered_by": s["triggered_by"],
            # "Operator/system": a person for a manual run or a second look, the
            # schedule otherwise.
            "operator": s["triggered_by_name"] or ("Schedule" if s["triggered_by"] == "SCHEDULE" else "System"),
            "scheduled_for": s["scheduled_for"], "started_at": s["started_at"], "launched_at": s["launched_at"],
            "ended_at": s["ended_at"], "duration_seconds": duration, "flew": flew, "status": s["status"],
            "reason": s["blocked_reason"] or s["failure_reason"] or s["abort_reason"],
            "aborted_by": s["aborted_by_name"], "distance_m": _f(s["distance_m"]),
            "planned": snap.get("estimate"),
            "battery_start_pct": track[0]["battery_pct"] if track else None,
            "battery_end_pct": track[-1]["battery_pct"] if track else None,
        },
        "route": {
            "base": base, "return_to_base": bool(route.get("return_to_base", True)),
            "waypoints": [{**dict(w), "altitude_m": _f(w["altitude_m"])} for w in waypoints],
            "waypoints_reached": sum(1 for w in waypoints if w["status"] in ("REACHED", "OBSERVED")),
            "track": [[t["latitude"], t["longitude"]] for t in track], "track_samples": total,
        },
        "summary": {
            "detections": sum(e["detections"] for e in out_events), "events": len(out_events),
            "suspicious": sum(1 for e in out_events
                              if e["risk_level"] in SUSPICIOUS and e["status"] != "FALSE_POSITIVE"),
            "incidents": len(incidents),
            "false_positives": sum(1 for e in out_events if e["status"] == "FALSE_POSITIVE"),
            "by_risk": by_risk, "by_module": by_module,
            "media": sum(len(v) for v in media.values()),
        },
        "events": out_events,
    }


# ═════════════════════════════════════════════════════════════════════════════
# Formatting shared by the renderers
# ═════════════════════════════════════════════════════════════════════════════

class _Clock:
    """Report times, in the organisation's zone."""

    def __init__(self, zone: str):
        try:
            self.tz = ZoneInfo(zone)
        except Exception:
            self.tz = ZoneInfo("Asia/Singapore")

    def dt(self, value, fallback: str = "—") -> str:
        if value is None:
            return fallback
        if isinstance(value, str):
            value = datetime.fromisoformat(value)
        return value.astimezone(self.tz).strftime("%d %b %Y %H:%M:%S")

    def t(self, value, fallback: str = "—") -> str:
        if value is None:
            return fallback
        if isinstance(value, str):
            value = datetime.fromisoformat(value)
        return value.astimezone(self.tz).strftime("%H:%M:%S")


def _duration(seconds: float | None) -> str:
    if seconds is None:
        return "—"
    m, s = divmod(int(round(seconds)), 60)
    h, m = divmod(m, 60)
    return f"{h}h {m:02d}m {s:02d}s" if h else f"{m}m {s:02d}s"


def _distance(m: float | None) -> str:
    if m is None:
        return "—"
    return f"{m / 1000:.2f} km" if m >= 1000 else f"{m:.0f} m"


def _module(module_type: str | None) -> str:
    return MODULE_NAMES.get(module_type or "", _pretty(module_type))


def _what(e: dict) -> str:
    return _module(e["module_type"]) + (f" · {e['label']}" if e.get("label") else "")


def _started_by(m: dict) -> str:
    """The mission's "operator/system": the person for a manual run, the schedule otherwise."""
    if m["triggered_by"] == "MANUAL":
        return f"{m['operator']} (manual run)"
    if m["triggered_by"] == "VERIFY_REQUEST":
        return f"{m['operator']} (verification request)"
    return "Schedule"


def _where(e: dict) -> str:
    if e["latitude"] is None or e["longitude"] is None:
        return "No position reported"
    return f"{e['latitude']:.6f}, {e['longitude']:.6f}"


def _confidence(e: dict) -> str:
    return f"{round(e['ai_confidence'] * 100)}%" if e["ai_confidence"] is not None else "—"


def _risk(e: dict) -> str:
    return f"{e['risk_level']} ({e['risk_score']})" if e["risk_score"] is not None else e["risk_level"]


def _incident(e: dict) -> str:
    inc = e["incident"]
    if not inc:
        return "None"
    head = inc["ref"] or inc["title"]
    guard = f", guard {inc['dispatched_guard_name']}" if inc["dispatched_guard_name"] else ""
    return f"{head} — {_pretty(inc['status'])}{guard}"


def _video(v: dict) -> str:
    where = {"central": "held centrally", "both": "held centrally and at the site",
             "local": f"held at the site ({_pretty(v['sync_state']).lower()})"}.get(v["location"], v["location"])
    length = f", {v['duration_seconds']:.0f}s" if v["duration_seconds"] else ""
    digest = f", SHA-256 {v['checksum_sha256'][:16]}…" if v["checksum_sha256"] else ""
    return f"{_pretty(v['kind'])} {v['media_id'][:8]}{length} — {where}{digest}"


def _cctv(c: dict) -> str:
    how = "covers the spot" if c["method"] == "COVERAGE" else "nearby"
    dist = f", {c['distance_m']:.0f} m" if c["distance_m"] is not None else ""
    saw = (f"; saw it too ({c['detections']} detection(s))" if c["corroborates"]
           else "; saw nothing matching")
    rec = "; recording available" if c["has_recording"] else ""
    return f"{c['camera_name'] or 'Camera'} ({how}{dist}){saw}{rec}"


# ═════════════════════════════════════════════════════════════════════════════
# Reading stored pictures
# ═════════════════════════════════════════════════════════════════════════════

def _read_local(path: str) -> bytes | None:
    root = Path(settings.EVIDENCE_ROOT).resolve()
    target = (root / path).resolve()
    if root not in target.parents or not target.is_file():
        return None
    return target.read_bytes()


def _read_s3(path: str) -> bytes | None:  # pragma: no cover - needs an object store
    from app.core.object_store import _get_client
    try:
        return _get_client().get_object(Bucket=settings.S3_BUCKET, Key=path)["Body"].read()
    except Exception:
        return None


def stored_file_reader() -> FileReader:
    """Reads a stored file by its storage path, from wherever this deployment
    keeps evidence. Returns None for a file that is not there."""
    return _read_s3 if settings.STORAGE_BACKEND == "s3" else _read_local


# ═════════════════════════════════════════════════════════════════════════════
# PDF
# ═════════════════════════════════════════════════════════════════════════════

def _route_drawing(data: dict, width_mm: float = 180, height_mm: float = 95):
    """The planned route and the flown track, drawn to scale, north up.

    A drawing, not a map: no tiles are fetched, so the report builds with no
    network and nothing about the site leaves the system to make it."""
    from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
    from reportlab.lib import colors
    from reportlab.lib.units import mm

    route = data["route"]
    planned = [(w["latitude"], w["longitude"]) for w in route["waypoints"]]
    base = route["base"]
    if base:
        b = (base["latitude"], base["longitude"])
        planned = [b] + planned + ([b] if route["return_to_base"] and planned else [])
    track = [tuple(p) for p in route["track"]]
    events = [(e["latitude"], e["longitude"], e["risk_level"]) for e in data["events"]
              if e["latitude"] is not None and e["longitude"] is not None]
    points = planned + track + [(a, b) for a, b, _ in events]
    if len(points) < 2:
        return None

    w, h, pad = width_mm * mm, height_mm * mm, 8 * mm
    lat0 = sum(p[0] for p in points) / len(points)
    kx = 111_320.0 * math.cos(math.radians(lat0))   # metres per degree of longitude here
    ky = 110_540.0
    xs = [p[1] * kx for p in points]
    ys = [p[0] * ky for p in points]
    span_x, span_y = (max(xs) - min(xs)) or 1.0, (max(ys) - min(ys)) or 1.0
    scale = min((w - 2 * pad) / span_x, (h - 2 * pad) / span_y)
    off_x = (w - span_x * scale) / 2 - min(xs) * scale
    off_y = (h - span_y * scale) / 2 - min(ys) * scale

    def xy(lat, lng):
        return lng * kx * scale + off_x, lat * ky * scale + off_y

    d = Drawing(w, h)
    d.add(Rect(0, 0, w, h, fillColor=colors.HexColor("#F7F8FA"), strokeColor=colors.HexColor("#CCCCCC"),
               strokeWidth=0.5))
    if len(planned) > 1:
        d.add(PolyLine([c for p in planned for c in xy(*p)], strokeColor=colors.HexColor("#4F6CFF"),
                       strokeWidth=1, strokeDashArray=[4, 3]))
    if len(track) > 1:
        d.add(PolyLine([c for p in track for c in xy(*p)], strokeColor=colors.HexColor("#0B8FAC"), strokeWidth=1.6))
    for wp in route["waypoints"]:
        x, y = xy(wp["latitude"], wp["longitude"])
        reached = wp["status"] in ("REACHED", "OBSERVED")
        d.add(Circle(x, y, 5, fillColor=colors.HexColor("#4F6CFF") if reached else colors.white,
                     strokeColor=colors.HexColor("#4F6CFF"), strokeWidth=1))
        d.add(String(x, y - 2.4, str(wp["sequence"]), fontSize=6.5, fontName="Helvetica-Bold", textAnchor="middle",
                     fillColor=colors.white if reached else colors.HexColor("#4F6CFF")))
    if base:
        x, y = xy(base["latitude"], base["longitude"])
        d.add(Rect(x - 5, y - 5, 10, 10, fillColor=colors.HexColor("#00A878"), strokeColor=colors.white,
                   strokeWidth=0.8))
        d.add(String(x, y - 2.4, "H", fontSize=6.5, fontName="Helvetica-Bold", textAnchor="middle",
                     fillColor=colors.white))
    # Rings, not dots: an event is usually at a waypoint, and must not hide it.
    for lat, lng, level in events:
        x, y = xy(lat, lng)
        d.add(Circle(x, y, 8, fillColor=None, strokeColor=colors.HexColor(RISK_HEX.get(level, "#888888")),
                     strokeWidth=1.8))
    # A scale bar at a round length, and north.
    bar_m = next((b for b in (5, 10, 20, 50, 100, 200, 500, 1000, 2000, 5000)
                  if b * scale >= 18 * mm), 5000)
    d.add(Line(pad, pad / 2, pad + bar_m * scale, pad / 2, strokeColor=colors.HexColor("#555555"), strokeWidth=1))
    d.add(String(pad + bar_m * scale + 3, pad / 2 - 2, _distance(bar_m), fontSize=6.5, fontName="Helvetica",
                 fillColor=colors.HexColor("#555555")))
    # North, drawn rather than typed: the standard PDF fonts have no arrow glyph.
    nx, ny = w - pad, h - pad - 10
    d.add(Line(nx, ny, nx, ny + 9, strokeColor=colors.HexColor("#555555"), strokeWidth=1))
    d.add(Line(nx, ny + 9, nx - 2.5, ny + 5, strokeColor=colors.HexColor("#555555"), strokeWidth=1))
    d.add(Line(nx, ny + 9, nx + 2.5, ny + 5, strokeColor=colors.HexColor("#555555"), strokeWidth=1))
    d.add(String(nx, ny - 8, "N", fontSize=7, fontName="Helvetica", textAnchor="middle",
                 fillColor=colors.HexColor("#555555")))
    return d


def _snapshot_flowable(e: dict, read_file: FileReader, styles):
    """The stored picture, or a plain statement of why there is none."""
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph

    snap = e["snapshot"]
    if snap["state"] == "none":
        return Paragraph("<b>No snapshot.</b> None was recorded for this event.", styles["BodyText"])
    if snap["state"] == "held_at_site":
        return Paragraph(
            f"<b>Snapshot held at the site.</b> The recording policy keeps it on the site's gateway "
            f"({escape(_pretty(snap.get('sync_state')).lower())}); it has not been uploaded.", styles["BodyText"])
    payload = read_file(snap["path"])
    if payload is None:
        return Paragraph("<b>Snapshot missing from storage.</b> The event recorded a picture but the file "
                         "is no longer present.", styles["BodyText"])
    try:
        img = Image(io.BytesIO(payload))
        ratio = img.imageHeight / float(img.imageWidth or 1)
        width = MAX_IMAGE_WIDTH_MM * mm
        height = width * ratio
        if height > MAX_IMAGE_HEIGHT_MM * mm:
            height = MAX_IMAGE_HEIGHT_MM * mm
            width = height / (ratio or 1)
        img.drawWidth, img.drawHeight = width, height
        return img
    except Exception:
        return Paragraph("<b>Snapshot unreadable.</b> The stored file exists but is not an image this report "
                         "can show.", styles["BodyText"])


def render_pdf(data: dict, read_file: FileReader | None = None) -> bytes:
    """The PDF, from a loaded report. `read_file` fetches a stored picture by
    its storage path; without it every picture is reported as missing."""
    _check_reportlab()
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    read_file = read_file or (lambda _path: None)
    m, clock = data["mission"], _Clock(data["timezone"])
    styles = getSampleStyleSheet()
    title = ParagraphStyle("drTitle", parent=styles["Title"], fontSize=16, spaceAfter=2)
    sub = ParagraphStyle("drSub", parent=styles["Normal"], fontSize=9, textColor=colors.HexColor("#555555"))
    cell = ParagraphStyle("drCell", parent=styles["BodyText"], fontSize=8, leading=10)
    grey, light = colors.HexColor("#555555"), colors.HexColor("#DDDDDD")

    def p(value) -> Paragraph:
        return Paragraph(escape(str(value)) if value not in (None, "") else "—", cell)

    def facts(rows: list[list], widths) -> Table:
        t = Table([[p(c) for c in r] for r in rows], colWidths=widths)
        t.setStyle(TableStyle([
            ("TEXTCOLOR", (0, 0), (0, -1), grey), ("TEXTCOLOR", (2, 0), (2, -1), grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
            ("TOPPADDING", (0, 0), (-1, -1), 1),
        ]))
        return t

    def grid(rows: list[list], widths, tints: dict[int, str] | None = None) -> Table:
        t = Table([[c if hasattr(c, "wrap") else p(c) for c in r] for r in rows], colWidths=widths, repeatRows=1)
        style = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F0F0F0")),
                 ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CCCCCC")),
                 ("VALIGN", (0, 0), (-1, -1), "TOP")]
        for row, hexcolor in (tints or {}).items():
            style.append(("BACKGROUND", (0, row), (0, row), colors.HexColor(hexcolor)))
        t.setStyle(TableStyle(style))
        return t

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm,
                            bottomMargin=15 * mm, title=f"Drone Patrol {m['session_number']}")
    story = [
        Paragraph("SEVENTH AI VISION", title),
        Paragraph("Drone Patrol Report", styles["Heading2"]),
        Paragraph(escape(m["session_number"]), sub),
        Spacer(1, 6), HRFlowable(width="100%", color=colors.HexColor("#999999")), Spacer(1, 8),
    ]

    # ── Mission ──
    drone = " · ".join(x for x in (m["drone_name"], m["drone_code"]) if x) or "None assigned"
    story += [facts([
        ["Site", m["site_name"], "Mission", m["mission_name"]],
        ["Drone", drone, "Started by", _started_by(m)],
        ["Date", clock.dt(m["scheduled_for"] or m["started_at"])[:11], "Status", _pretty(m["status"])],
        ["Start", clock.dt(m["launched_at"] or m["started_at"], "Did not start"),
         "End", clock.dt(m["ended_at"], "Not ended")],
        ["Duration", _duration(m["duration_seconds"]) if m["flew"] else "Did not fly",
         "Distance", _distance(m["distance_m"])],
        ["Route", m["route_name"], "AI profile", m["profile_name"] or "Defaults"],
    ], [24 * mm, 66 * mm, 24 * mm, 66 * mm]), Spacer(1, 4)]
    if m["reason"]:
        label = {"BLOCKED": "Blocked before launch", "MISSED": "Missed", "ABORTED": "Aborted",
                 "FAILED": "Failed", "CANCELLED": "Cancelled"}.get(m["status"], "Note")
        who = f" (by {escape(m['aborted_by'])})" if m["aborted_by"] else ""
        story += [Paragraph(f"<b>{label}{who}:</b> {escape(m['reason'])}", styles["BodyText"]), Spacer(1, 4)]
    story += [Paragraph(f"All times are {escape(data['timezone'])}.", sub), Spacer(1, 8)]

    # ── Route ──
    route = data["route"]
    story += [Paragraph("Route", styles["Heading3"])]
    drawing = _route_drawing(data)
    if drawing is not None:
        flown = (f"flown track from {route['track_samples']} position samples (solid)" if route["track"]
                 else "no flown track was recorded")
        story += [drawing, Paragraph(
            f"Planned route (dashed) and {flown}. Numbered circles are waypoints, filled when reached; "
            f"H is the launch point; coloured rings are events.", sub), Spacer(1, 6)]
    if route["waypoints"]:
        rows = [["#", "Waypoint", "Position", "Status", "Reached", "Left"]]
        for w in route["waypoints"]:
            rows.append([w["sequence"], w["name"] or "—", f"{w['latitude']:.6f}, {w['longitude']:.6f}",
                         _pretty(w["status"]), clock.t(w["reached_at"]), clock.t(w["departed_at"])])
        story += [grid(rows, [8 * mm, 46 * mm, 46 * mm, 26 * mm, 27 * mm, 27 * mm]),
                  Paragraph(f"{route['waypoints_reached']} of {len(route['waypoints'])} waypoints reached.", sub)]
    else:
        story += [Paragraph("This flight had no waypoints.", sub)]
    story += [Spacer(1, 10)]

    # ── AI summary ──
    sm = data["summary"]
    block = [Paragraph("AI summary", styles["Heading3"]), facts([
        ["Detections", sm["detections"], "Events", sm["events"]],
        ["Suspicious events", sm["suspicious"], "Incidents", sm["incidents"]],
        ["False positives", sm["false_positives"], "Snapshots and clips", sm["media"]],
    ], [34 * mm, 56 * mm, 34 * mm, 56 * mm]), Spacer(1, 4)]
    if sm["events"]:
        levels = grid([["Risk level", "Events"]] + [[level, sm["by_risk"].get(level, 0)] for level in RISK_LEVELS],
                      [40 * mm, 25 * mm],
                      {i: "#" + RISK_TINT[level] for i, level in enumerate(RISK_LEVELS, start=1)})
        levels.hAlign = "LEFT"
        block += [levels,
                  Paragraph("Suspicious: assessed MEDIUM or above and not marked a false positive. "
                            "Risk is how serious an event is here and now; AI confidence, shown separately "
                            "below, is how sure the AI is of what it saw.", sub)]
    else:
        block += [Paragraph("The drone raised no events on this flight.", styles["BodyText"])]
    # One block: a table of five rows split across two pages reads as two tables.
    story += [KeepTogether(block), Spacer(1, 10)]

    # ── Event details ──
    if data["events"]:
        story += [Paragraph("Event details", styles["Heading3"])]
    for e in data["events"]:
        place = " · ".join(x for x in (e["zone_name"] and f"{e['zone_name']} ({_pretty(e['zone_type']).lower()})",
                                       _where(e)) if x)
        seen = f"{e['detections']}×" + (f" over {e['observed_seconds']:.0f}s" if e["observed_seconds"] else "")
        head = [
            HRFlowable(width="100%", color=light), Spacer(1, 3),
            Paragraph(f"{e['number']}. {escape(_what(e))} "
                      f"<font color='{RISK_HEX.get(e['risk_level'], '#555555')}'>— {escape(_risk(e))}</font>",
                      styles["Heading4"]),
            facts([
                ["Time", clock.dt(e["detected_at"]), "Waypoint",
                 e["waypoint_sequence"] if e["waypoint_sequence"] is not None else "Between waypoints"],
                ["Location", place, "Seen", seen],
                ["AI confidence", _confidence(e), "Risk", _risk(e)],
                ["Verification", _pretty(e["verification_state"]), "Incident", _incident(e)],
                ["Officer action", e["officer_action"], "Status", _pretty(e["status"])],
            ], [24 * mm, 66 * mm, 24 * mm, 66 * mm]),
            Spacer(1, 4),
        ]
        body = [_snapshot_flowable(e, read_file, styles), Spacer(1, 4)]
        if e["risk_factors"]:
            why = "; ".join(f"{escape(str(f.get('detail') or f.get('factor')))} "
                            f"({'+' if (f.get('points') or 0) > 0 else ''}{f.get('points')})"
                            for f in e["risk_factors"])
            body += [Paragraph(f"<b>Why this risk:</b> {why}", cell)]
        videos = "; ".join(escape(_video(v)) for v in e["videos"]) or "None recorded"
        body += [Paragraph(f"<b>Video:</b> {videos}", cell)]
        cctv = "; ".join(escape(_cctv(c)) for c in e["cctv"]) or "No fixed camera covers or is near this spot"
        body += [Paragraph(f"<b>Related CCTV:</b> {cctv}", cell), Spacer(1, 10)]
        # An event is one block: its picture and what was done never part from
        # its heading at a page break.
        story += [KeepTogether(head + body)]

    footer = (f"Drone patrol {m['session_number']} · generated {clock.dt(data['generated_at'])} "
              f"{data['timezone']} from the flight's own record")

    def on_page(canvas, document):
        # On every page, so a page pulled out of the stack still says what it is.
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(grey)
        canvas.drawString(15 * mm, 9 * mm, footer)
        canvas.drawRightString(A4[0] - 15 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=on_page, onLaterPages=on_page)
    return buf.getvalue()


# ═════════════════════════════════════════════════════════════════════════════
# Excel
# ═════════════════════════════════════════════════════════════════════════════

def _workbook():
    _check_openpyxl()
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    header_font, header_fill = Font(bold=True), PatternFill("solid", fgColor="EEEEEE")

    def sheet(title: str, headers: list[str], widths: list[int], first: bool = False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = header_font, header_fill
        ws.freeze_panes = "A2"
        for i, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        return ws

    return wb, sheet


def _risk_fill(level: str):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=RISK_TINT.get(level, "FFFFFF"))


def _save(wb) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def render_xlsx(data: dict) -> bytes:
    """Five sheets: the flight, its waypoints, its events, the evidence behind
    them and the fixed cameras that looked at the same place. No pictures — a
    spreadsheet is for sorting and totalling; the picture is in the PDF, and the
    checksum here says which file it was."""
    wb, sheet = _workbook()
    m, sm, clock = data["mission"], data["summary"], _Clock(data["timezone"])

    ws = sheet("Summary", ["Field", "Value"], [26, 52], first=True)
    for label, value in (
        ("Session", m["session_number"]), ("Site", m["site_name"]), ("Mission", m["mission_name"]),
        ("Drone", " · ".join(x for x in (m["drone_name"], m["drone_code"]) if x) or "None assigned"),
        ("Route", m["route_name"] or ""), ("AI profile", m["profile_name"] or "Defaults"),
        ("Started by", _started_by(m)),
        ("Scheduled for", clock.dt(m["scheduled_for"], "")),
        ("Start", clock.dt(m["launched_at"] or m["started_at"], "Did not start")),
        ("End", clock.dt(m["ended_at"], "Not ended")),
        ("Duration", _duration(m["duration_seconds"]) if m["flew"] else "Did not fly"),
        ("Status", _pretty(m["status"])), ("Reason", m["reason"] or ""),
        ("Distance flown", _distance(m["distance_m"])),
        ("Waypoints reached", f"{data['route']['waypoints_reached']} of {len(data['route']['waypoints'])}"),
        ("Detections", sm["detections"]), ("Events", sm["events"]), ("Suspicious events", sm["suspicious"]),
        ("Incidents", sm["incidents"]), ("False positives", sm["false_positives"]),
        *[(f"Events at {level}", sm["by_risk"].get(level, 0)) for level in RISK_LEVELS],
        ("Times are", data["timezone"]), ("Generated", clock.dt(data["generated_at"])),
    ):
        ws.append([label, value])

    ws = sheet("Waypoints", ["#", "Name", "Latitude", "Longitude", "Altitude (m)", "Hover (s)", "Observe (s)",
                             "Snapshot", "Status", "Reached", "Left"], [5, 26, 12, 12, 12, 10, 11, 10, 12, 20, 20])
    for w in data["route"]["waypoints"]:
        ws.append([w["sequence"], w["name"] or "", w["latitude"], w["longitude"], w["altitude_m"],
                   w["hover_seconds"], w["observe_seconds"], "Yes" if w["snapshot_required"] else "",
                   _pretty(w["status"]), clock.dt(w["reached_at"], ""), clock.dt(w["departed_at"], "")])

    ws = sheet("Events", ["#", "Time", "Detection", "Label", "Zone", "Latitude", "Longitude", "Waypoint",
                          "AI confidence", "Risk level", "Risk score", "Detections", "Seen (s)", "Verification",
                          "Status", "Officer action", "Incident", "Incident status", "Guard dispatched"],
               [5, 20, 16, 16, 22, 12, 12, 10, 14, 12, 11, 11, 10, 14, 16, 40, 18, 16, 22])
    for e in data["events"]:
        inc = e["incident"] or {}
        ws.append([e["number"], clock.dt(e["detected_at"]), _module(e["module_type"]), e["label"] or "",
                   e["zone_name"] or "", e["latitude"], e["longitude"], e["waypoint_sequence"],
                   # A fraction with a percent format, so it sorts and averages as a number.
                   e["ai_confidence"], e["risk_level"], e["risk_score"], e["detections"], e["observed_seconds"],
                   _pretty(e["verification_state"]), _pretty(e["status"]), e["officer_action"],
                   inc.get("ref") or inc.get("title") or "", _pretty(inc.get("status")) if inc else "",
                   inc.get("dispatched_guard_name") or ""])
        ws.cell(row=ws.max_row, column=9).number_format = "0%"
        ws.cell(row=ws.max_row, column=10).fill = _risk_fill(e["risk_level"])
    if not data["events"]:
        ws.append(["The drone raised no events on this flight."])

    ws = sheet("Evidence", ["Event #", "Kind", "Captured", "Where it is", "Seconds", "Bytes", "SHA-256", "Reference"],
               [9, 14, 20, 34, 9, 12, 66, 38])
    for e in data["events"]:
        snap = e["snapshot"]
        if snap["state"] != "none":
            where = {"available": "Held centrally" + (" (AI worker evidence)" if snap["source"] == "ai_worker" else ""),
                     "held_at_site": "Held at the site"}[snap["state"]]
            ws.append([e["number"], "Snapshot", clock.dt(snap["captured_at"], ""), where, None, None,
                       snap["checksum_sha256"] or "", snap["path"] or ""])
        for v in e["videos"]:
            where = {"central": "Held centrally", "both": "Held centrally and at the site",
                     "local": f"Held at the site ({_pretty(v['sync_state']).lower()})"}.get(v["location"], v["location"])
            ws.append([e["number"], _pretty(v["kind"]), clock.dt(v["captured_at"], ""), where,
                       v["duration_seconds"], v["size_bytes"], v["checksum_sha256"] or "", v["media_id"]])

    ws = sheet("CCTV", ["Event #", "Camera", "How", "Distance (m)", "Saw it too", "Detections", "Recording"],
               [9, 30, 18, 13, 12, 12, 12])
    for e in data["events"]:
        for c in e["cctv"]:
            ws.append([e["number"], c["camera_name"] or "", "Covers the spot" if c["method"] == "COVERAGE" else "Nearby",
                       c["distance_m"], "Yes" if c["corroborates"] else "No", c["detections"],
                       "Yes" if c["has_recording"] else "No"])
    return _save(wb)


async def build_pdf(db: AsyncSession, session_id) -> bytes:
    data = await load_report(db, session_id)
    return await asyncio.to_thread(render_pdf, data, stored_file_reader())


async def build_xlsx(db: AsyncSession, session_id) -> bytes:
    return render_xlsx(await load_report(db, session_id))


# ═════════════════════════════════════════════════════════════════════════════
# Storing the finished report
# ═════════════════════════════════════════════════════════════════════════════

def _write(rel: str, payload: bytes, content_type: str) -> None:
    if settings.STORAGE_BACKEND == "s3":  # pragma: no cover - needs an object store
        from app.core.object_store import put_object
        put_object(rel, payload, content_type)
        return
    path = Path(settings.EVIDENCE_ROOT) / rel
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_bytes(payload)
    tmp.replace(path)


async def store_reports(db: AsyncSession, session_id) -> list[dict]:
    """Write the PDF and the workbook to storage and record them against the
    flight, each with its SHA-256. Idempotent per (flight, format): a retry
    replaces the row rather than adding a second claim to a second report.

    A format that cannot be rendered is logged and left out; the other is still
    stored, and the download endpoint reports the real failure when asked."""
    data = await load_report(db, session_id)
    day = data["mission"]["ended_at"] or data["mission"]["started_at"] or datetime.now(timezone.utc)
    prefix = f"drone/{data['tenant_id']}/reports/{day:%Y/%m/%d}/{data['session_id']}"
    reader = stored_file_reader()
    stored: list[dict] = []
    for fmt, build, ext, mime in (
        ("PDF", lambda: render_pdf(data, reader), "pdf", "application/pdf"),
        ("XLSX", lambda: render_xlsx(data), "xlsx",
         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    ):
        try:
            payload = await asyncio.to_thread(build)
            rel = f"{prefix}/report.{ext}"
            await asyncio.to_thread(_write, rel, payload, mime)
        except Exception:
            logger.exception("drone report: could not produce the %s report for %s", fmt, session_id)
            continue
        # Hashed from the bytes just written, not read back: reading back would
        # hash whatever landed, and record a truncated write as correct.
        row = (await db.execute(text("""
            INSERT INTO drone_reports (tenant_id, session_id, report_format, storage_path, file_bytes, checksum_sha256)
            VALUES (CAST(:t AS uuid), CAST(:s AS uuid), :f, :p, :n, :sum)
            ON CONFLICT (session_id, report_format) DO UPDATE
                SET storage_path = EXCLUDED.storage_path, file_bytes = EXCLUDED.file_bytes,
                    checksum_sha256 = EXCLUDED.checksum_sha256, generated_at = now()
            RETURNING id, report_format, storage_path, file_bytes, checksum_sha256, generated_at
        """), {"t": data["tenant_id"], "s": data["session_id"], "f": fmt, "p": rel, "n": len(payload),
               "sum": hashlib.sha256(payload).hexdigest()})).mappings().first()
        stored.append(dict(row))
    return stored


# ═════════════════════════════════════════════════════════════════════════════
# A period: the daily, weekly and monthly summary
# ═════════════════════════════════════════════════════════════════════════════

def period_bounds(start: date, end: date, tz: ZoneInfo) -> tuple[datetime, datetime]:
    """[start 00:00, the day after end 00:00) in the organisation's zone."""
    return (datetime.combine(start, time.min, tzinfo=tz),
            datetime.combine(end + timedelta(days=1), time.min, tzinfo=tz))


async def load_period(db: AsyncSession, *, start: date, end: date, site_id=None, mission_id=None,
                      allowed_site_ids: list[str] | None = None) -> dict:
    """Every flight created in the period, and the events they raised, for one
    scope: the organisation, a site or a mission. `allowed_site_ids` narrows it
    to the sites a caller may see; None means no restriction."""
    tz = await tenant_tz(db)
    since, until = period_bounds(start, end, tz)
    params = {"since": since, "until": until, "site": str(site_id) if site_id else None,
              "mission": str(mission_id) if mission_id else None, "allowed": allowed_site_ids}
    where = """
           AND (CAST(:site AS uuid) IS NULL OR ps.site_id = CAST(:site AS uuid))
           AND (CAST(:mission AS uuid) IS NULL OR ps.mission_id = CAST(:mission AS uuid))
           AND (CAST(:allowed AS text[]) IS NULL OR ps.site_id::text = ANY(CAST(:allowed AS text[])))"""
    sessions = (await db.execute(text(f"""
        SELECT ps.id, ps.session_number, ps.mission_name, ps.drone_name, ps.status, ps.triggered_by,
               ps.scheduled_for, ps.started_at, ps.launched_at, ps.ended_at, ps.distance_m, ps.event_count,
               ps.incident_count, ps.created_at,
               COALESCE(ps.blocked_reason, ps.failure_reason, ps.abort_reason) AS reason, si.name AS site_name
          FROM drone_patrol_sessions ps LEFT JOIN sites si ON si.id = ps.site_id
         WHERE ps.created_at >= :since AND ps.created_at < :until {where}
         ORDER BY ps.created_at
    """), params)).mappings().all()
    events = (await db.execute(text(f"""
        SELECT e.detected_at, e.module_type, e.label, e.zone_name, e.ai_confidence, e.risk_level, e.risk_score,
               e.status, e.verification_state, e.incident_id, ps.session_number, ps.mission_name,
               si.name AS site_name
          FROM drone_events e
          JOIN drone_patrol_sessions ps ON ps.id = e.session_id
          LEFT JOIN sites si ON si.id = ps.site_id
         WHERE ps.created_at >= :since AND ps.created_at < :until {where}
         ORDER BY e.detected_at
    """), params)).mappings().all()

    by_status: dict[str, int] = {}
    flown_s, metres = 0.0, 0.0
    for s in sessions:
        by_status[s["status"]] = by_status.get(s["status"], 0) + 1
        if s["launched_at"] and s["ended_at"]:
            flown_s += (s["ended_at"] - s["launched_at"]).total_seconds()
        metres += float(s["distance_m"] or 0)
    by_risk = {level: 0 for level in RISK_LEVELS}
    for e in events:
        by_risk[e["risk_level"]] = by_risk.get(e["risk_level"], 0) + 1
    return {
        "start": start, "end": end, "timezone": tz.key, "generated_at": datetime.now(timezone.utc),
        "flights": [dict(s) for s in sessions],
        "events": [{**dict(e), "ai_confidence": _f(e["ai_confidence"])} for e in events],
        "totals": {
            "flights": len(sessions), "by_status": by_status,
            "completed": by_status.get("COMPLETED", 0),
            "did_not_complete": sum(by_status.get(k, 0) for k in ("FAILED", "ABORTED", "BLOCKED", "MISSED")),
            "flight_seconds": flown_s, "distance_m": metres,
            "events": len(events), "by_risk": by_risk,
            "suspicious": sum(1 for e in events if e["risk_level"] in SUSPICIOUS and e["status"] != "FALSE_POSITIVE"),
            "false_positives": sum(1 for e in events if e["status"] == "FALSE_POSITIVE"),
            "incidents": len({e["incident_id"] for e in events if e["incident_id"]}),
        },
    }


def render_period_xlsx(data: dict, *, scope_label: str) -> bytes:
    """Summary, a row per flight, a row per event."""
    wb, sheet = _workbook()
    t, clock = data["totals"], _Clock(data["timezone"])

    ws = sheet("Summary", ["Field", "Value"], [28, 44], first=True)
    for label, value in (
        ("Covers", scope_label),
        ("Period", f"{data['start']:%d %b %Y} – {data['end']:%d %b %Y}"),
        ("Flights", t["flights"]), ("Completed", t["completed"]),
        ("Did not complete (failed, aborted, blocked, missed)", t["did_not_complete"]),
        *[(f"Flights {_pretty(k).lower()}", v) for k, v in sorted(t["by_status"].items())],
        ("Time in the air", _duration(t["flight_seconds"])), ("Distance flown", _distance(t["distance_m"])),
        ("Events", t["events"]), ("Suspicious events", t["suspicious"]), ("Incidents", t["incidents"]),
        ("False positives", t["false_positives"]),
        *[(f"Events at {level}", t["by_risk"].get(level, 0)) for level in RISK_LEVELS],
        ("Times are", data["timezone"]), ("Generated", clock.dt(data["generated_at"])),
    ):
        ws.append([label, value])

    ws = sheet("Flights", ["Session", "Site", "Mission", "Drone", "Started by", "Status", "Scheduled for", "Start",
                           "End", "Minutes", "Distance (m)", "Events", "Incidents", "Reason"],
               [24, 22, 24, 18, 14, 13, 20, 20, 20, 9, 12, 8, 10, 50])
    for s in data["flights"]:
        minutes = (round((s["ended_at"] - s["launched_at"]).total_seconds() / 60, 1)
                   if s["launched_at"] and s["ended_at"] else None)
        ws.append([s["session_number"], s["site_name"] or "", s["mission_name"] or "", s["drone_name"] or "",
                   _pretty(s["triggered_by"]), _pretty(s["status"]), clock.dt(s["scheduled_for"], ""),
                   clock.dt(s["launched_at"] or s["started_at"], ""), clock.dt(s["ended_at"], ""), minutes,
                   _f(s["distance_m"]), s["event_count"], s["incident_count"], s["reason"] or ""])

    ws = sheet("Events", ["Time", "Site", "Mission", "Session", "Detection", "Label", "Zone", "AI confidence",
                          "Risk level", "Risk score", "Verification", "Status", "Incident"],
               [20, 22, 24, 24, 16, 16, 22, 14, 12, 11, 14, 16, 10])
    for e in data["events"]:
        ws.append([clock.dt(e["detected_at"]), e["site_name"] or "", e["mission_name"] or "", e["session_number"],
                   _module(e["module_type"]), e["label"] or "", e["zone_name"] or "", e["ai_confidence"],
                   e["risk_level"], e["risk_score"], _pretty(e["verification_state"]), _pretty(e["status"]),
                   "Yes" if e["incident_id"] else ""])
        ws.cell(row=ws.max_row, column=8).number_format = "0%"
        ws.cell(row=ws.max_row, column=9).fill = _risk_fill(e["risk_level"])
    if not data["events"]:
        ws.append(["No events in this period."])
    return _save(wb)


def period_text(data: dict, *, scope_label: str) -> str:
    """The email body for a summary: the numbers someone reads without opening
    the attachment."""
    t = data["totals"]
    risks = ", ".join(f"{t['by_risk'][level]} {level.lower()}" for level in RISK_LEVELS if t["by_risk"].get(level))
    lines = [
        f"Drone patrol summary for {scope_label}, {data['start']:%d %b %Y} to {data['end']:%d %b %Y}.",
        "",
        f"Flights: {t['flights']} ({t['completed']} completed, {t['did_not_complete']} did not complete).",
        f"Time in the air: {_duration(t['flight_seconds'])}. Distance: {_distance(t['distance_m'])}.",
        f"Events: {t['events']}" + (f" ({risks})." if risks else "."),
        f"Suspicious events: {t['suspicious']}. Incidents: {t['incidents']}. False positives: {t['false_positives']}.",
        "",
        "The attached workbook lists every flight and every event in the period.",
    ]
    return "\n".join(lines)


def flight_text(data: dict) -> str:
    """The email body for one flight's report."""
    m, sm, clock = data["mission"], data["summary"], _Clock(data["timezone"])
    lines = [
        f"Drone patrol {m['session_number']}: {m['mission_name'] or 'mission'} at {m['site_name'] or 'site'}.",
        "",
        f"Status: {_pretty(m['status'])}" + (f" — {m['reason']}" if m["reason"] else "") + ".",
        f"Start: {clock.dt(m['launched_at'] or m['started_at'], 'did not start')}. "
        f"End: {clock.dt(m['ended_at'], 'not ended')} ({data['timezone']}).",
        f"Events: {sm['events']} ({sm['suspicious']} suspicious). Incidents: {sm['incidents']}.",
        "",
        "The attached report has the route, every event with its snapshot, the related CCTV and what was done.",
    ]
    return "\n".join(lines)
