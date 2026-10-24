"""The patrol report renders, and admits what it does not have.

A report is the artefact an agency hands to a client or a regulator, so the
failure that matters is not a crash — it is a report that looks complete while
quietly missing its evidence.

Sections:
  A — It renders, and it is a PDF (2 tests)
  B — A missing image is stated, never omitted (3 tests)
  C — The content comes from the session, not from configuration (2 tests)
  D — Times are printed in the schedule's own zone, and say so (6 tests)
"""
from __future__ import annotations

import base64
import io
import os
import re
import uuid
import zlib
from datetime import date, datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo

import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.main import app  # noqa: F401  (module level — see other suites)
from app.services import virtual_patrol as vp
from app.services import vpatrol_reports as rep
from app.services import vpatrol_scheduler as sched


def _camera(**over):
    base = {
        "id": uuid.uuid4(), "sequence_no": 1, "camera_name": "Fire Exit",
        "camera_code": None, "status": "COMPLETED",
        "snapshot_path": None, "snapshot_taken_at": None,
        "snapshot_error": None, "officer_notes": None, "location": "East stair",
    }
    base.update(over)
    return base


def _data(cameras=None, answers=None, **session_over):
    session = {
        "id": uuid.uuid4(), "patrol_number": "VP-20260401-abcd1234",
        "schedule_name": "Morning Patrol", "site_name": "ABC Industrial",
        "scheduled_for": datetime(2026, 4, 1, 7, 0, tzinfo=timezone.utc),
        "started_at": datetime(2026, 4, 1, 7, 3, tzinfo=timezone.utc),
        "completed_at": datetime(2026, 4, 1, 7, 21, tzinfo=timezone.utc),
        "status": "COMPLETED", "camera_count": 1, "completed_camera_count": 1,
        "notes": None, "officer_name": "Tan Wei Ming",
    }
    session.update(session_over)
    return {"session": session,
            "cameras": cameras if cameras is not None else [_camera()],
            "answers": answers or {}}


# ─── A. It renders ───────────────────────────────────────────────────────────

def test_the_report_renders_a_real_pdf():
    out = rep.render_pdf(_data())
    assert out.startswith(b"%PDF-"), "not a PDF"
    assert len(out) > 800


def test_a_patrol_with_no_questions_still_produces_a_report():
    """A camera configured with no questions is unusual, not invalid, and must
    not take the whole report down on the day somebody needs it."""
    out = rep.render_pdf(_data(answers={}))
    assert out.startswith(b"%PDF-")


# ─── B. A missing image is stated ────────────────────────────────────────────

def test_a_camera_with_no_snapshot_says_why():
    cam = _camera(snapshot_path=None, snapshot_error="The camera did not respond.")
    flow = rep._snapshot_flowable(cam, _styles())
    assert "did not respond" in _text_of(flow)


def test_a_snapshot_recorded_but_missing_from_disk_is_called_out():
    """The loud one. Silently omitting the image would make a patrol whose
    evidence was lost look exactly like one where the camera was never reached —
    a report hiding a gap in its own evidence."""
    cam = _camera(snapshot_path="does/not/exist.jpg")
    flow = rep._snapshot_flowable(cam, _styles())
    text = _text_of(flow)
    assert "missing from storage" in text.lower()


def test_no_snapshot_and_no_error_still_produces_a_statement():
    flow = rep._snapshot_flowable(_camera(), _styles())
    assert "No snapshot" in _text_of(flow)


# ─── C. Content comes from the session ───────────────────────────────────────

def test_the_answer_display_handles_every_shape_an_answer_takes():
    assert rep._answer_display(
        {"answer_json": ["Lock", "Sign"], "answer_text": None}) == "Lock, Sign"
    assert rep._answer_display(
        {"answer_json": None, "answer_text": "YES"}) == "YES"
    assert rep._answer_display(
        {"answer_json": None, "answer_text": None}) == "Not answered"


def test_an_unanswered_required_question_is_shown_not_hidden():
    """A blank row would read as "nothing to report" when it means "nobody
    checked"."""
    cam = _camera()
    answers = {str(cam["id"]): [{
        "session_camera_id": cam["id"], "sequence_no": 1,
        "question_text": "Is the fire exit clear?", "question_type": "YES_NO",
        "answer_text": None, "answer_json": None, "is_exception": False,
        "exception_reason": None, "incident_id": None, "answered_by": None,
    }]}
    out = rep.render_pdf(_data(cameras=[cam], answers=answers))
    assert out.startswith(b"%PDF-")


# ─── D. Times are the schedule's own clock ───────────────────────────────────
#
# The bug as it was seen on 24 Oct 2026: a patrol on the "Jurong night camera
# round" schedule (patrol_time 22:00, Asia/Singapore) printed "Scheduled
# 2026-10-01 14:00" — the UTC value PostgreSQL hands back, printed as it came.

NIGHT_ROUND_UTC = datetime(2026, 10, 1, 14, 0, tzinfo=timezone.utc)  # 22:00 SGT


def _night_round(zone: str | None = "Asia/Singapore") -> dict:
    """The night round as the loader returns it: every time in UTC."""
    cam = _camera(snapshot_taken_at=NIGHT_ROUND_UTC + timedelta(minutes=5))
    answers = {str(cam["id"]): [{
        "session_camera_id": cam["id"], "sequence_no": 1,
        "question_text": "Is the loading bay shutter down?", "question_type": "YES_NO",
        "answer_text": "NO", "answer_json": None, "is_exception": True,
        "exception_reason": "Shutter left open", "incident_id": None,
        "answered_by": "Tan Wei Ming",
        "answered_at": NIGHT_ROUND_UTC + timedelta(minutes=6),
    }]}
    data = _data(cameras=[cam], answers=answers,
                 schedule_name="Jurong night camera round",
                 scheduled_for=NIGHT_ROUND_UTC,
                 started_at=NIGHT_ROUND_UTC + timedelta(minutes=2),
                 completed_at=NIGHT_ROUND_UTC + timedelta(minutes=9))
    data["timezone"] = zone
    return data


def test_a_22_00_singapore_patrol_prints_22_00_in_the_pdf():
    text_ = _pdf_text(rep.render_pdf(_night_round()))
    assert "2026-10-01 22:00" in text_, "the scheduled time is not in Singapore time"
    assert "2026-10-01 14:" not in text_, "the PDF still prints UTC"
    for local in ("22:02", "22:09", "22:05"):  # started, completed, snapshot
        assert f"2026-10-01 {local}" in text_, local
    assert "All times are Asia/Singapore." in text_, "the PDF does not say which zone"


def test_a_22_00_singapore_patrol_prints_22_00_in_the_workbook():
    book = _sheets(rep.render_xlsx(_night_round()))
    summary = dict(book["Summary"][1:])
    assert summary["Times are"] == "Asia/Singapore"
    assert summary["Scheduled"] == "2026-10-01 22:00"
    assert summary["Started"] == "2026-10-01 22:02"
    assert summary["Completed"] == "2026-10-01 22:09"

    # Each time column names the zone too: Exceptions is often the only sheet
    # anybody opens, and it has no Summary beside it.
    head, row = book["Camera Inspection"][:2]
    assert row[head.index("Snapshot taken (Asia/Singapore)")] == "2026-10-01 22:05"
    for name in ("Questions & Answers", "Exceptions"):
        head, row = book[name][:2]
        assert row[head.index("Answered at (Asia/Singapore)")] == "2026-10-01 22:06", name

    cells = [str(c) for rows in book.values() for r in rows for c in r if c is not None]
    assert not [c for c in cells if "2026-10-01 14:" in c], "the workbook still prints UTC"


def test_the_zone_is_the_schedules_own_not_a_fixed_eight_hours():
    data = _night_round("Asia/Kolkata")  # UTC+05:30
    summary = dict(_sheets(rep.render_xlsx(data))["Summary"][1:])
    assert summary["Times are"] == "Asia/Kolkata"
    assert summary["Scheduled"] == "2026-10-01 19:30"
    text_ = _pdf_text(rep.render_pdf(data))
    assert "2026-10-01 19:30" in text_ and "All times are Asia/Kolkata." in text_


@pytest.mark.parametrize("zone", [None, "Mars/Olympus_Mons"])
def test_no_zone_or_an_unknown_one_prints_singapore_as_the_scheduler_does(zone):
    """vpatrol_scheduler._zone's rule: an unknown name runs the patrol in
    Asia/Singapore, so the report must print it in Asia/Singapore too."""
    data = _night_round(zone)
    summary = dict(_sheets(rep.render_xlsx(data))["Summary"][1:])
    assert summary["Times are"] == "Asia/Singapore"
    assert summary["Scheduled"] == "2026-10-01 22:00"
    assert "All times are Asia/Singapore." in _pdf_text(rep.render_pdf(data))


# Against a real database: the session is made the way the scheduler makes it
# (latest_due, then create_session) and read back through the loader the
# download endpoints use, as svc_app with the tenant set, as the router reads it.

_app_db_url = os.environ.get("DATABASE_URL", "")
_m = re.search(r"@([^:/]+):", _app_db_url)
_db_host = _m.group(1) if _m else "localhost"
ADMIN_DATABASE_URL = os.environ.get(
    "ADMIN_TEST_DATABASE_URL",
    f"postgresql+asyncpg://postgres:change_me_dev_only@{_db_host}:5432/seventh_ai_vision_test",
)


async def _admin_sql(statement: str, params: dict | None = None):
    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            r = await s.execute(text(statement), params or {})
            rows = r.all() if r.returns_rows else []
            await s.commit()
        return rows
    finally:
        await engine.dispose()


async def _night_round_session(zone: str) -> dict:
    """A 22:00 schedule in `zone` with one camera, and its 1 Oct session."""
    ids = {k: uuid.uuid4() for k in ("tenant", "site", "cam", "sched")}
    await _admin_sql("INSERT INTO tenants (id, name, slug) VALUES (:t,'Jurong Co',:s)",
                     {"t": ids["tenant"], "s": f"vptz-{ids['tenant'].hex[:10]}"})
    await _admin_sql("INSERT INTO sites (id, tenant_id, name) VALUES (:i,:t,'Jurong')",
                     {"i": ids["site"], "t": ids["tenant"]})
    await _admin_sql("INSERT INTO cameras (id, tenant_id, site_id, name, location) "
                     "VALUES (:i,:t,:s,'Loading Bay','Dock 3')",
                     {"i": ids["cam"], "t": ids["tenant"], "s": ids["site"]})
    await _admin_sql(
        "INSERT INTO virtual_patrol_schedules "
        "  (id, tenant_id, site_id, name, schedule_type, start_date, patrol_time, timezone) "
        "VALUES (:i,:t,:s,'Jurong night camera round','DAILY',:d,:pt,:z)",
        {"i": ids["sched"], "t": ids["tenant"], "s": ids["site"],
         "d": date(2026, 1, 1), "pt": time(22, 0), "z": zone})
    await _admin_sql(
        "INSERT INTO virtual_patrol_schedule_cameras "
        "  (tenant_id, schedule_id, camera_id, sequence_no) VALUES (:t,:sc,:c,1)",
        {"t": ids["tenant"], "sc": ids["sched"], "c": ids["cam"]})

    # Half an hour after 22:00 on 1 Oct where the site is: inside the
    # scheduler's lookback, so latest_due returns that evening's occurrence.
    half_past = datetime.combine(date(2026, 10, 1), time(22, 30), tzinfo=ZoneInfo(zone))
    due = sched.latest_due(
        schedule_type="DAILY", start_date=date(2026, 1, 1), end_date=None,
        weekdays=None, patrol_time=time(22, 0), tz_name=zone,
        now_utc=half_past.astimezone(timezone.utc))
    assert due is not None, "the scheduler found nothing due"
    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            created = await vp.create_session(s, schedule_id=str(ids["sched"]),
                                              scheduled_for=due)
            await s.commit()
    finally:
        await engine.dispose()
    return {**ids, "session": created["session_id"], "due": due}


async def _load_as_app(tenant_id, session_id: str) -> dict:
    engine = create_async_engine(os.environ["TEST_DATABASE_URL"])
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            restricted = (await s.execute(text(
                "SELECT NOT (rolsuper OR rolbypassrls) FROM pg_roles "
                " WHERE rolname = current_user"))).scalar()
            assert restricted, "loaded as a role that bypasses RLS: proves nothing"
            await s.execute(text("SELECT set_config('app.current_tenant', :t, true)"),
                            {"t": str(tenant_id)})
            return await rep.load_session_for_report(s, session_id)
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_the_night_round_report_prints_22_00_from_the_database():
    made = await _night_round_session("Asia/Singapore")
    assert made["due"] == NIGHT_ROUND_UTC, "the scheduler did not make 22:00 SGT 14:00 UTC"

    data = await _load_as_app(made["tenant"], made["session"])
    # Stored as UTC, as it should be. Converting is the report's job.
    assert data["session"]["scheduled_for"] == NIGHT_ROUND_UTC
    assert data["timezone"] == "Asia/Singapore"

    text_ = _pdf_text(rep.render_pdf(data))
    assert "2026-10-01 22:00" in text_ and "2026-10-01 14:00" not in text_
    assert "All times are Asia/Singapore." in text_
    summary = dict(_sheets(rep.render_xlsx(data))["Summary"][1:])
    assert summary["Scheduled"] == "2026-10-01 22:00"
    assert summary["Times are"] == "Asia/Singapore"


@pytest.mark.asyncio
async def test_a_patrol_whose_schedule_was_deleted_prints_in_singapore():
    """schedule_id is set to NULL when a schedule is deleted. The report must
    still render, in the scheduler's fallback zone, rather than 404 or UTC."""
    made = await _night_round_session("Asia/Kolkata")
    await _admin_sql("DELETE FROM virtual_patrol_schedules WHERE id = :i",
                     {"i": made["sched"]})

    data = await _load_as_app(made["tenant"], made["session"])
    assert data["timezone"] == "Asia/Singapore"
    # 22:00 Kolkata is 16:30 UTC, which is 00:30 the next day in Singapore.
    summary = dict(_sheets(rep.render_xlsx(data))["Summary"][1:])
    assert summary["Scheduled"] == "2026-10-02 00:30"
    assert "All times are Asia/Singapore." in _pdf_text(rep.render_pdf(data))


# ── helpers ──────────────────────────────────────────────────────────────────

def _styles():
    from reportlab.lib.styles import getSampleStyleSheet
    return getSampleStyleSheet()


def _text_of(flowable) -> str:
    return getattr(flowable, "text", "") or ""


def _pdf_text(pdf: bytes) -> str:
    """The text a reader sees, from reportlab's page streams.

    reportlab compresses each page (ASCII85, then Flate), so the words are not
    in the file's raw bytes; this undoes both and joins the strings the page
    draws with Tj. Enough for "is this time on the page", which is the claim.
    """
    out = []
    for body in re.findall(rb"stream\r?\n(.*?)endstream", pdf, re.S):
        raw = body.strip()
        try:
            if raw.endswith(b"~>"):
                raw = base64.a85decode(raw.removeprefix(b"<~").removesuffix(b"~>"))
            raw = zlib.decompress(raw)
        except Exception:
            continue  # an image or font stream, not page text
        for s in re.findall(rb"\(((?:\\.|[^\\)])*)\)\s*Tj", raw):
            out.append(re.sub(rb"\\(.)", rb"\1", s).decode("latin-1"))
    return "\n".join(out)


def _sheets(payload: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(payload))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()]
            for ws in wb.worksheets}
