"""Digests: which window, queued once, and what the summary says.

The period tests are the point of this file, and they are pure functions — no
database, no clock. A digest that silently covers the wrong week still arrives
looking authoritative, and nobody checks the dates on an email that says what
they expect. So the boundaries are pinned explicitly: the Monday of last week,
the last day of last month, and a February that does not exist in every year.

Sections:
  A — Which window a frequency means (8 tests)
  B — Queued once, for a closed window, across tenants (4 tests)
  C — What the digest contains (3 tests)
  D — Times are the schedule's own clock (5 tests)
"""
from __future__ import annotations

import os
import re
import uuid
from datetime import date, datetime, time, timedelta, timezone

import pytest
from sqlalchemy import text
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

# Module level on purpose — app.main pulls the ML stack, and inside a test that
# import lands on whichever test runs first and trips pytest-timeout.
from app.main import app  # noqa: F401
from app.services import vpatrol_digest as digest
from app.services import vpatrol_email as mail

_app_db_url = os.environ.get("DATABASE_URL", "")
_m = re.search(r"@([^:/]+):", _app_db_url)
_db_host = _m.group(1) if _m else "localhost"
ADMIN_DATABASE_URL = os.environ.get(
    "ADMIN_TEST_DATABASE_URL",
    f"postgresql+asyncpg://postgres:change_me_dev_only@{_db_host}:5432/seventh_ai_vision_test",
)


async def _sql(stmt: str, params: dict | None = None):
    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            r = await s.execute(text(stmt), params or {})
            rows = r.mappings().all() if r.returns_rows else []
            await s.commit()
        return rows
    finally:
        await engine.dispose()


# ─── A. Which window a frequency means ───────────────────────────────────────

def test_daily_means_yesterday():
    got = mail.digest_period("DAILY", today=date(2026, 9, 12))
    assert got == (date(2026, 9, 11), date(2026, 9, 11))


def test_weekly_means_last_monday_to_sunday():
    """2026-09-12 is a Saturday. Last week is Mon 31 Aug – Sun 6 Sep, not the
    seven days before today — a rolling seven days would overlap the week
    already reported and double-count every patrol on the seam."""
    got = mail.digest_period("WEEKLY", today=date(2026, 9, 12))
    assert got == (date(2026, 8, 31), date(2026, 9, 6))
    assert got[0].isoweekday() == 1 and got[1].isoweekday() == 7


def test_weekly_run_on_a_monday_reports_the_week_that_just_ended():
    """The sharpest boundary. On Monday the window must be the week that closed
    yesterday, and must not include today."""
    monday = date(2026, 9, 7)
    assert monday.isoweekday() == 1
    start, end = mail.digest_period("WEEKLY", today=monday)
    assert (start, end) == (date(2026, 8, 31), date(2026, 9, 6))
    assert end < monday


def test_monthly_means_the_previous_calendar_month():
    got = mail.digest_period("MONTHLY", today=date(2026, 9, 12))
    assert got == (date(2026, 8, 1), date(2026, 8, 31))


def test_monthly_on_the_first_reports_the_month_that_just_ended():
    got = mail.digest_period("MONTHLY", today=date(2026, 9, 1))
    assert got == (date(2026, 8, 1), date(2026, 8, 31))


def test_monthly_handles_a_short_february():
    """2026 is not a leap year, so February ends on the 28th. Arithmetic that
    assumes 30 or 31 days silently drops or invents a day of patrols."""
    got = mail.digest_period("MONTHLY", today=date(2026, 3, 15))
    assert got == (date(2026, 2, 1), date(2026, 2, 28))


def test_monthly_across_a_year_boundary():
    got = mail.digest_period("MONTHLY", today=date(2026, 1, 8))
    assert got == (date(2025, 12, 1), date(2025, 12, 31))


def test_immediate_and_nonsense_have_no_window():
    """IMMEDIATE has no window by definition. An unrecognised value must not
    fall through to a default — a typo in a tenant's configuration becoming
    "daily" is a change nobody asked for."""
    assert mail.digest_period("IMMEDIATE", today=date(2026, 9, 12)) is None
    assert mail.digest_period("FORTNIGHTLY", today=date(2026, 9, 12)) is None


# ─── B. Queued once, for a closed window ─────────────────────────────────────

async def _schedule_with_patrols(*, frequency: str, recipients: list[str],
                                 patrol_days: list[date], tz: str = "UTC"):
    ids = {k: uuid.uuid4() for k in ("tenant", "site", "sched")}
    await _sql("INSERT INTO tenants (id, name, slug) VALUES (:t,'Digest Co',:s)",
               {"t": ids["tenant"], "s": f"vpdig-{ids['tenant'].hex[:10]}"})
    await _sql("INSERT INTO sites (id, tenant_id, name) VALUES (:i,:t,'Site')",
               {"i": ids["site"], "t": ids["tenant"]})
    await _sql(
        "INSERT INTO virtual_patrol_schedules "
        "  (id, tenant_id, site_id, name, schedule_type, start_date, patrol_time, "
        "   timezone, email_frequency) "
        "VALUES (:i,:t,:s,'Morning Patrol','DAILY',:d,:pt,:tz,:f)",
        {"i": ids["sched"], "t": ids["tenant"], "s": ids["site"],
         "d": date(2026, 1, 1), "pt": time(7, 0), "tz": tz, "f": frequency})
    for email in recipients:
        await _sql("INSERT INTO virtual_patrol_email_recipients "
                   "  (tenant_id, schedule_id, email) VALUES (:t,:s,:e)",
                   {"t": ids["tenant"], "s": ids["sched"], "e": email})
    for n, day in enumerate(patrol_days):
        # A DISTINCT TIME PER PATROL, even on the same day.
        #
        # uq_vpsess_execution is UNIQUE (schedule_id, scheduled_for) -- one
        # execution, one session -- so two patrols at 07:00 on the same day are
        # not something the application permits, and seeding them is invalid
        # test data rather than a scenario worth covering. Hours keep them
        # inside the same local date for the window comparison.
        await _sql(
            "INSERT INTO virtual_patrol_sessions "
            "  (tenant_id, site_id, schedule_id, patrol_number, schedule_name, "
            "   scheduled_for, status, camera_count, completed_camera_count) "
            "VALUES (:t,:s,:sc,:n,'Morning Patrol',:w,'COMPLETED',3,3)",
            {"t": ids["tenant"], "s": ids["site"], "sc": ids["sched"],
             "n": f"VPD-{ids['sched'].hex[:6]}-{n}",
             "w": datetime.combine(day, time(6 + n, 0), tzinfo=timezone.utc)})
    return ids


async def _run(now: datetime):
    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            return await mail.enqueue_due_digests(s, now=now)
    finally:
        await engine.dispose()


@pytest.mark.asyncio
async def test_a_daily_digest_is_queued_for_yesterdays_patrols():
    today = date(2026, 9, 12)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[today - timedelta(days=1)])
    await _run(datetime.combine(today, time(6, 0), tzinfo=timezone.utc))

    rows = await _sql("SELECT frequency, period_start, period_end, recipients, status "
                      "  FROM virtual_patrol_email_queue WHERE schedule_id = :s",
                      {"s": ids["sched"]})
    assert len(rows) == 1, rows
    assert rows[0]["frequency"] == "DAILY"
    assert rows[0]["period_start"] == today - timedelta(days=1)
    assert rows[0]["status"] == "PENDING"


@pytest.mark.asyncio
async def test_running_twice_does_not_queue_the_same_digest_again():
    """The heart of it. This runs every two minutes; a digest queued per tick
    would send the same weekly summary 720 times a day."""
    today = date(2026, 9, 12)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[today - timedelta(days=1)])
    now = datetime.combine(today, time(6, 0), tzinfo=timezone.utc)
    await _run(now)
    second = await _run(now)

    rows = await _sql("SELECT count(*) AS n FROM virtual_patrol_email_queue "
                      " WHERE schedule_id = :s", {"s": ids["sched"]})
    assert rows[0]["n"] == 1, f"queued {rows[0]['n']} digests for one window"
    assert second["already_queued"] >= 1


@pytest.mark.asyncio
async def test_a_window_with_no_patrols_queues_nothing():
    """An agency that receives "0 patrols" every Monday stops reading Monday's
    email, and then misses the week something did go wrong."""
    today = date(2026, 9, 12)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[today - timedelta(days=30)])   # outside yesterday
    await _run(datetime.combine(today, time(6, 0), tzinfo=timezone.utc))

    rows = await _sql("SELECT count(*) AS n FROM virtual_patrol_email_queue "
                      " WHERE schedule_id = :s", {"s": ids["sched"]})
    assert rows[0]["n"] == 0


@pytest.mark.asyncio
async def test_a_schedule_with_no_recipients_queues_nothing():
    today = date(2026, 9, 12)
    ids = await _schedule_with_patrols(
        frequency="WEEKLY", recipients=[],
        patrol_days=[today - timedelta(days=8)])
    await _run(datetime.combine(today, time(6, 0), tzinfo=timezone.utc))

    rows = await _sql("SELECT count(*) AS n FROM virtual_patrol_email_queue "
                      " WHERE schedule_id = :s", {"s": ids["sched"]})
    assert rows[0]["n"] == 0


# ─── C. What the digest contains ─────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_summary_counts_what_happened():
    today = date(2026, 9, 12)
    yesterday = today - timedelta(days=1)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[yesterday, yesterday])

    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            data = await digest.load_digest(s, schedule_id=str(ids["sched"]),
                                            start=yesterday, end=yesterday,
                                            tz_name="UTC")
    finally:
        await engine.dispose()

    text_body = digest.summary_text(data)
    assert "Patrols scheduled:   2" in text_body, text_body
    assert "completed:           2" in text_body, text_body


@pytest.mark.asyncio
async def test_the_workbook_has_a_row_per_patrol():
    today = date(2026, 9, 12)
    yesterday = today - timedelta(days=1)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[yesterday, yesterday, yesterday])

    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            payload, body = await digest.build(s, schedule_id=str(ids["sched"]),
                                               start=yesterday, end=yesterday,
                                               tz_name="UTC")
    finally:
        await engine.dispose()

    assert payload[:2] == b"PK", "that is not an xlsx"
    import io as _io

    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(payload))
    ws = wb["Patrols"]
    # Title, dates, blank, header, then one row per patrol.
    assert ws.max_row == 4 + 3, f"expected 3 patrol rows, sheet has {ws.max_row} rows"
    assert "Findings" in wb.sheetnames


@pytest.mark.asyncio
async def test_no_findings_is_stated_rather_than_left_blank():
    """An empty sheet reads as a broken export, and "no findings" is the result
    an agency most wants to be able to show a client."""
    today = date(2026, 9, 12)
    yesterday = today - timedelta(days=1)
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"],
        patrol_days=[yesterday])

    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            payload, _ = await digest.build(s, schedule_id=str(ids["sched"]),
                                            start=yesterday, end=yesterday,
                                            tz_name="UTC")
    finally:
        await engine.dispose()

    import io as _io

    from openpyxl import load_workbook
    ws = load_workbook(_io.BytesIO(payload))["Findings"]
    cells = [c.value for row in ws.iter_rows() for c in row if c.value]
    assert any("No findings" in str(v) for v in cells), cells


# ─── D. Times are the schedule's own clock ───────────────────────────────────
#
# The same bug the patrol report had until 24 Oct 2026: PostgreSQL hands back
# UTC, and the digest printed it as it came, so a finding answered at 22:00 in
# Singapore read "14:00" in the email and in its workbook.

NIGHT_ROUND_UTC = datetime(2026, 10, 1, 14, 0, tzinfo=timezone.utc)  # 22:00 SGT


def _night_digest(zone: str | None = "Asia/Singapore") -> dict:
    """A digest as load_digest returns it: every time in UTC. `...` leaves the
    zone out altogether, as data built by an older caller would."""
    data = {
        "schedule": {"name": "Jurong night camera round", "site_name": "Jurong"},
        "start": date(2026, 10, 1), "end": date(2026, 10, 1),
        "patrols": [{"patrol_number": "VP-0042", "scheduled_for": NIGHT_ROUND_UTC,
                     "officer_name": "Tan Wei Ming", "status": "COMPLETED",
                     "camera_count": 1, "completed_camera_count": 1,
                     "exception_count": 1}],
        "exceptions": [{"patrol_number": "VP-0042", "camera_name": "Loading bay",
                        "question_text": "Is the shutter down?", "answer_text": "NO",
                        "answered_at": NIGHT_ROUND_UTC, "incident_id": None}],
    }
    if zone is not ...:
        data["timezone"] = zone
    return data


def _sheets(payload: bytes) -> dict[str, list[tuple]]:
    import io as _io

    from openpyxl import load_workbook
    wb = load_workbook(_io.BytesIO(payload))
    return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)]
            for ws in wb.worksheets}


def _cell(rows: list[tuple], header_row: int, column: str):
    head = rows[header_row]
    return rows[header_row + 1][head.index(column)]


def test_a_finding_answered_at_22_00_singapore_prints_22_00_in_the_email():
    body = digest.summary_text(_night_digest())
    assert "01 Oct 22:00  Loading bay" in body, body
    assert "14:00" not in body, "the email still prints UTC"
    assert "All times are Asia/Singapore." in body, "the email does not say which zone"


def test_a_finding_answered_at_22_00_singapore_prints_22_00_in_the_workbook():
    book = _sheets(digest.render_xlsx(_night_digest()))
    # Patrols: title, dates, blank, then the header on row 4.
    assert _cell(book["Patrols"], 3, "Scheduled (Asia/Singapore)") == "2026-10-01 22:00"
    assert _cell(book["Findings"], 0, "When (Asia/Singapore)") == "2026-10-01 22:00"
    cells = [str(c) for rows in book.values() for r in rows for c in r if c is not None]
    assert not [c for c in cells if "14:00" in c], "the workbook still prints UTC"


def test_the_zone_is_the_schedules_own_not_a_fixed_eight_hours():
    data = _night_digest("Asia/Kolkata")  # UTC+05:30
    body = digest.summary_text(data)
    assert "01 Oct 19:30" in body and "All times are Asia/Kolkata." in body, body
    book = _sheets(digest.render_xlsx(data))
    assert _cell(book["Findings"], 0, "When (Asia/Kolkata)") == "2026-10-01 19:30"


@pytest.mark.parametrize("zone", [..., None, "Mars/Olympus_Mons"])
def test_no_zone_or_an_unknown_one_prints_singapore_as_the_scheduler_does(zone):
    """vpatrol_scheduler._zone's rule: an unknown name runs the patrol in
    Asia/Singapore, so the digest must print it in Asia/Singapore too."""
    data = _night_digest(zone)
    body = digest.summary_text(data)
    assert "01 Oct 22:00" in body and "All times are Asia/Singapore." in body, body


@pytest.mark.asyncio
async def test_the_digest_built_from_the_database_prints_singapore_time():
    """End to end through build(), as vpatrol_email calls it: the zone the
    window was read in is the zone the times are printed in."""
    ids = await _schedule_with_patrols(
        frequency="DAILY", recipients=["ops@example.test"], patrol_days=[],
        tz="Asia/Singapore")
    session, scam, question = (uuid.uuid4() for _ in range(3))
    await _sql(
        "INSERT INTO virtual_patrol_sessions "
        "  (id, tenant_id, site_id, schedule_id, patrol_number, schedule_name, "
        "   scheduled_for, status, camera_count, completed_camera_count) "
        "VALUES (:i,:t,:s,:sc,:n,'Morning Patrol',:w,'COMPLETED',1,1)",
        {"i": session, "t": ids["tenant"], "s": ids["site"], "sc": ids["sched"],
         "n": f"VPD-{ids['sched'].hex[:6]}-night", "w": NIGHT_ROUND_UTC})
    await _sql(
        "INSERT INTO virtual_patrol_session_cameras "
        "  (id, tenant_id, session_id, sequence_no, camera_name, status) "
        "VALUES (:i,:t,:se,1,'Loading bay','COMPLETED')",
        {"i": scam, "t": ids["tenant"], "se": session})
    await _sql(
        "INSERT INTO virtual_patrol_session_questions "
        "  (id, tenant_id, session_camera_id, question_text, question_type, "
        "   is_required, sequence_no) "
        "VALUES (:i,:t,:c,'Is the shutter down?','YES_NO',TRUE,1)",
        {"i": question, "t": ids["tenant"], "c": scam})
    await _sql(
        "INSERT INTO virtual_patrol_session_answers "
        "  (tenant_id, session_question_id, answer_text, answered_at, is_exception) "
        "VALUES (:t,:q,'NO',:at,TRUE)",
        {"t": ids["tenant"], "q": question, "at": NIGHT_ROUND_UTC})

    engine = create_async_engine(ADMIN_DATABASE_URL)
    factory = async_sessionmaker(engine, expire_on_commit=False, class_=AsyncSession)
    try:
        async with factory() as s:
            payload, body = await digest.build(
                s, schedule_id=str(ids["sched"]), start=date(2026, 10, 1),
                end=date(2026, 10, 1), tz_name="Asia/Singapore")
    finally:
        await engine.dispose()

    assert "Findings raised:     1" in body, body
    assert "01 Oct 22:00  Loading bay" in body and "14:00" not in body, body
    assert "All times are Asia/Singapore." in body, body
    book = _sheets(payload)
    assert _cell(book["Patrols"], 3, "Scheduled (Asia/Singapore)") == "2026-10-01 22:00"
    assert _cell(book["Findings"], 0, "When (Asia/Singapore)") == "2026-10-01 22:00"
