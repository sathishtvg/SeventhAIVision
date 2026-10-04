"""Drone patrol, phase 11: the patrol report and its delivery.

  A — What the report says: the flight's own record, the stored picture, the
      same numbers in the PDF's data and the workbook
  B — Storing: once per flight and format, with a checksum that matches the file
  C — Emails: who gets the immediate report, retries and backoff, summaries per
      scope and closed period in the organisation's own day
  D — The job across tenants, as the application's database user
  E — The API: who may read, export and choose recipients

The report job is called directly with a chosen clock and a mail sender that
records instead of sending, so every email here is the real document, built from
the real rows, and no mail server is involved.
"""
from __future__ import annotations

import hashlib
import io
import uuid
from datetime import date, datetime, time, timedelta, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from sqlalchemy import text

# Module level on purpose: app.main pulls the ML stack.
from app.main import app  # noqa: F401
from app.core.config import settings
from app.db.session import AsyncSessionLocal
from app.services import drone_report_delivery as delivery
from app.services import drone_reports as reports
from app.services import drone_runner as runner
from tests.test_drone_ai_pipeline import _ai, _detect, _events, _flying, _last, _world
from tests.test_drone_api import ADMIN, GUARD, OPERATOR, SUPERVISOR, VIEWER
from tests.test_drone_api import _world as _roles_world
from tests.test_drone_edge_sync import _client, _run, _sql

SGT = ZoneInfo("Asia/Singapore")
XLSX = "vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class Outbox:
    """Stands in for the mail server: keeps what it was handed, or refuses."""

    def __init__(self):
        self.sent: list[dict] = []
        self.down = False

    async def __call__(self, recipients, subject, body, filename, payload, subtype):
        if self.down:
            raise RuntimeError("Connection refused by mail server")
        self.sent.append({"to": recipients, "subject": subject, "body": body, "filename": filename,
                          "payload": payload, "subtype": subtype})


@pytest.fixture
def evidence(tmp_path, monkeypatch) -> Path:
    """Reports and pictures go to a directory of the test's own."""
    monkeypatch.setattr(settings, "EVIDENCE_ROOT", str(tmp_path))
    return tmp_path


def _jpeg() -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (64, 36), (40, 90, 160)).save(buf, "JPEG")
    return buf.getvalue()


async def _flight_with_event(w: dict, at: datetime) -> tuple[str, dict]:
    """A flown flight with one verified HIGH intrusion and its incident, ended
    four minutes after the sighting."""
    sid = await _flying(w, at)
    for s in range(3):
        await _detect(w, at + timedelta(seconds=2 * s), "intrusion")
    await _ai(at + timedelta(seconds=10))
    await _ai(at + timedelta(seconds=12))
    await _sql("UPDATE drone_patrol_sessions SET status = 'COMPLETED', ended_at = :e, distance_m = 412.5 WHERE id = :s",
               {"e": at + timedelta(minutes=4), "s": uuid.UUID(sid)})
    [event] = await _events(sid)
    return sid, event


async def _snapshot_for(w: dict, event: dict, evidence: Path, *, location: str = "central",
                        write: bool = True) -> str:
    rel = f"drone/{w['tenant']}/2026/snap-{uuid.uuid4().hex}.jpg"
    payload = _jpeg()
    if write:
        (evidence / rel).parent.mkdir(parents=True, exist_ok=True)
        (evidence / rel).write_bytes(payload)
    await _sql("INSERT INTO drone_event_media (tenant_id, event_id, session_id, media_kind, storage_path, "
               "    storage_location, sync_state, checksum_sha256, size_bytes, captured_at) "
               "VALUES (:t,:e,:s,'SNAPSHOT',:p,:loc,:sync,:sum,:n,:at)",
               {"t": w["tenant"], "e": event["id"], "s": event["session_id"], "p": rel, "loc": location,
                "sync": "pending" if location == "local" else "synced",
                "sum": hashlib.sha256(payload).hexdigest(), "n": len(payload), "at": event["detected_at"]})
    return rel


async def _recipient(w: dict, email: str, frequency: str = "IMMEDIATE", *, site=None, mission=None,
                     active: bool = True, created: datetime | None = None) -> None:
    await _sql("INSERT INTO drone_report_recipients (tenant_id, site_id, mission_id, email, frequency, is_active, "
               "    created_at) VALUES (:t,:s,:m,:e,:f,:a,:c)",
               {"t": w["tenant"], "s": site, "m": mission, "e": email, "f": frequency, "a": active,
                "c": created or datetime.now(timezone.utc) - timedelta(days=30)})


async def _queue(w: dict) -> list[dict]:
    return [dict(r) for r in await _sql(
        "SELECT * FROM drone_report_email_queue WHERE tenant_id = :t ORDER BY created_at", {"t": w["tenant"]})]


async def _loaded(w: dict, sid: str) -> dict:
    async with AsyncSessionLocal() as db:
        await db.execute(text("SELECT set_config('app.current_tenant', :t, true)"), {"t": str(w["tenant"])})
        return await reports.load_report(db, sid)


async def _tick(w: dict, now: datetime, outbox: Outbox) -> dict:
    return await runner.process_reports(AsyncSessionLocal, str(w["tenant"]), now, outbox)


def _sheets(payload: bytes) -> dict:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(payload))
    return {ws.title: [[c.value for c in row] for row in ws.iter_rows()] for ws in wb.worksheets}


# ─── Pure ────────────────────────────────────────────────────────────────────

def test_an_address_is_normalised_or_refused():
    assert delivery.normalise_email("  Ops.Lead@Example.COM ") == "ops.lead@example.com"
    for bad in ("", "no-at-sign", "two@@example.com", "a@b", "a b@example.com", "a@example.com,b@example.com",
                "a@example.com\nBcc: x@example.com", "<a@example.com>"):
        assert delivery.normalise_email(bad) is None, bad


def test_a_scope_is_one_thing():
    s, m = uuid.uuid4(), uuid.uuid4()
    assert delivery.scope_key(None, None) == "tenant"
    assert delivery.scope_key(s, None) == f"site:{s}"
    assert delivery.scope_key(None, m) == f"mission:{m}"


def test_a_period_is_whole_days_in_the_organisations_zone():
    since, until = reports.period_bounds(date(2026, 9, 24), date(2026, 9, 24), SGT)
    assert since == datetime(2026, 9, 23, 16, 0, tzinfo=timezone.utc)
    assert until == datetime(2026, 9, 24, 16, 0, tzinfo=timezone.utc)


def test_what_an_officer_did_is_said_in_a_sentence():
    base = {"acknowledged_by_name": "Mei", "resolved_by_name": "Ravi", "false_positive_reason": "A guard on rounds"}
    assert reports.officer_action({**base, "status": "NEW"}) == "No action recorded"
    assert reports.officer_action({**base, "status": "ACKNOWLEDGED"}) == "Acknowledged by Mei"
    assert reports.officer_action({**base, "status": "RESOLVED"}) == "Resolved by Ravi"
    assert reports.officer_action({**base, "status": "FALSE_POSITIVE"}) == \
        "Marked a false positive by Ravi: A guard on rounds"


# ─── A. What the report says ─────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_report_is_the_flights_own_record(evidence):
    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    at = _last(2, 17)
    sid, event = await _flight_with_event(w, at)
    await _snapshot_for(w, event, evidence)

    data = await _loaded(w, sid)
    m, sm, [e] = data["mission"], data["summary"], data["events"]
    assert (m["site_name"], m["mission_name"], m["drone_name"], m["drone_code"]) == \
        ("Depot", "Night Watch", "Drone One", "D-01")
    assert m["operator"] == "Admin User" and m["triggered_by"] == "MANUAL"
    assert m["status"] == "COMPLETED" and m["flew"] is True and m["duration_seconds"] == 360.0
    assert m["distance_m"] == 412.5 and data["timezone"] == "Asia/Singapore"
    assert [wp["sequence"] for wp in data["route"]["waypoints"]] == [1]
    assert data["route"]["base"] == {"latitude": 1.3, "longitude": 103.8}
    assert data["route"]["track_samples"] == 36 and len(data["route"]["track"]) == 36

    assert sm["events"] == 1 and sm["suspicious"] == 1 and sm["incidents"] == 1 and sm["false_positives"] == 0
    assert sm["detections"] == 3 and sm["by_risk"]["HIGH"] == 1 and sm["by_module"] == {"intrusion": 1}
    assert e["module_type"] == "intrusion" and e["risk_level"] == "HIGH" and e["zone_name"] == "Loading Bay"
    # Two numbers, never one: how sure the AI is, and how much it matters.
    assert e["ai_confidence"] == 0.85 and e["risk_score"] == event["risk_score"]
    assert e["incident"]["ref"] and e["incident"]["status"] == "open"
    assert e["snapshot"]["state"] == "available" and e["snapshot"]["source"] == "drone"

    # Renaming the mission and the route afterwards does not rewrite history.
    await _run([("UPDATE drone_missions SET name = 'Renamed' WHERE id = :m", {"m": w["mission"]}),
                ("UPDATE drone_routes SET name = 'Rerouted' WHERE id = :r", {"r": w["route"]})])
    again = await _loaded(w, sid)
    assert again["mission"]["mission_name"] == "Night Watch" and again["mission"]["route_name"] == "Perimeter"


@pytest.mark.asyncio
async def test_the_pdf_shows_the_stored_picture_and_says_so_when_there_is_none(evidence):
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Image, Paragraph

    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    sid, event = await _flight_with_event(w, _last(2, 17))
    rel = await _snapshot_for(w, event, evidence)
    # Names are the customer's own text; markup in one must not break the page.
    await _sql("UPDATE drone_patrol_sessions SET mission_name = 'Gate <A> & yard' WHERE id = :s", {"s": uuid.UUID(sid)})
    data = await _loaded(w, sid)
    reader, styles = reports.stored_file_reader(), getSampleStyleSheet()

    pdf = reports.render_pdf(data, reader)
    assert pdf.startswith(b"%PDF") and b"/Subtype /Image" in pdf, "the stored snapshot is not in the PDF"
    assert isinstance(reports._snapshot_flowable(data["events"][0], reader, styles), Image)

    # The file gone from storage: said plainly, not left out.
    (evidence / rel).unlink()
    missing = reports._snapshot_flowable(data["events"][0], reader, styles)
    assert isinstance(missing, Paragraph) and "missing from storage" in missing.text
    assert b"/Subtype /Image" not in reports.render_pdf(data, reader)

    # A path that climbs out of the evidence store reads as "not there".
    assert reader("../../etc/passwd") is None


@pytest.mark.asyncio
async def test_a_picture_still_at_the_site_or_from_the_ai_worker_is_named_for_what_it_is(evidence):
    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    sid, event = await _flight_with_event(w, _last(2, 17))
    await _snapshot_for(w, event, evidence, location="local", write=False)
    held = (await _loaded(w, sid))["events"][0]["snapshot"]
    assert held["state"] == "held_at_site" and held["sync_state"] == "pending" and held["path"] is None

    # The AI worker's own evidence for one of the detections: the platform's
    # table, read only, used when the drone's upload is not at the centre.
    det = (await _sql("SELECT detection_id FROM drone_observations WHERE event_id = :e LIMIT 1",
                      {"e": event["id"]}))[0]["detection_id"]
    await _sql("INSERT INTO evidence (tenant_id, detection_id, media_type, storage_path, checksum_sha256, captured_at) "
               "VALUES (:t,:d,'image','evidence/worker.jpg','abc123',:at)",
               {"t": w["tenant"], "d": det, "at": event["detected_at"]})
    worker = (await _loaded(w, sid))["events"][0]["snapshot"]
    assert (worker["state"], worker["source"], worker["path"]) == ("available", "ai_worker", "evidence/worker.jpg")


@pytest.mark.asyncio
async def test_the_workbook_carries_the_same_numbers_as_the_report(evidence):
    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    sid, event = await _flight_with_event(w, _last(2, 17))
    await _snapshot_for(w, event, evidence)
    data = await _loaded(w, sid)
    book = _sheets(reports.render_xlsx(data))

    assert list(book) == ["Summary", "Waypoints", "Events", "Evidence", "CCTV"]
    summary = dict((r[0], r[1]) for r in book["Summary"][1:])
    assert summary["Mission"] == "Night Watch" and summary["Status"] == "Completed"
    assert summary["Events"] == 1 and summary["Suspicious events"] == 1 and summary["Incidents"] == 1
    assert summary["Events at HIGH"] == 1 and summary["Times are"] == "Asia/Singapore"
    header, row = book["Events"][0], book["Events"][1]
    cell = dict(zip(header, row))
    assert cell["Detection"] == "Intrusion" and cell["Zone"] == "Loading Bay"
    # A number, so it sorts and averages; and in its own column, apart from risk.
    assert cell["AI confidence"] == 0.85 and cell["Risk level"] == "HIGH" and cell["Risk score"] == event["risk_score"]
    assert cell["Incident"] == data["events"][0]["incident"]["ref"]
    evidence_row = dict(zip(book["Evidence"][0], book["Evidence"][1]))
    assert evidence_row["Kind"] == "Snapshot" and len(evidence_row["SHA-256"]) == 64


@pytest.mark.asyncio
async def test_a_flight_that_never_launched_still_has_a_report(evidence):
    w = await _world()
    await _sql("UPDATE drones SET battery_level = 4 WHERE id = :d", {"d": w["drone"]})
    async with _client() as c:
        r = await c.post(f"/api/v1/drone-missions/{w['mission']}/run", headers=w["h_admin"])
    assert r.json()["session"]["status"] == "BLOCKED", r.text
    data = await _loaded(w, r.json()["session"]["id"])
    m = data["mission"]
    assert m["status"] == "BLOCKED" and m["flew"] is False and m["reason"], m
    assert data["route"]["track"] == [] and data["summary"]["events"] == 0
    assert reports.render_pdf(data).startswith(b"%PDF")
    assert dict((r[0], r[1]) for r in _sheets(reports.render_xlsx(data))["Summary"][1:])["Duration"] == "Did not fly"


# ─── B. Storing ──────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_a_finished_flight_is_stored_once_with_a_checksum_that_matches_the_file(evidence):
    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    at = _last(2, 17)
    sid, _ = await _flight_with_event(w, at)
    ended, out = at + timedelta(minutes=4), Outbox()

    # Not the instant it lands: corroboration and an officer's first action
    # are still arriving.
    early = await _tick(w, ended + timedelta(minutes=1), out)
    assert early["stored"] == 0
    assert (await _sql("SELECT count(*) AS n FROM drone_reports WHERE session_id = :s", {"s": uuid.UUID(sid)}))[0]["n"] == 0

    got = await _tick(w, ended + timedelta(minutes=6), out)
    assert got["stored"] == 2, got
    rows = await _sql("SELECT report_format, storage_path, file_bytes, checksum_sha256 FROM drone_reports "
                      " WHERE session_id = :s ORDER BY report_format", {"s": uuid.UUID(sid)})
    assert [r["report_format"] for r in rows] == ["PDF", "XLSX"]
    for r in rows:
        on_disk = (evidence / r["storage_path"]).read_bytes()
        assert len(on_disk) == r["file_bytes"]
        assert hashlib.sha256(on_disk).hexdigest() == r["checksum_sha256"]
        assert str(w["tenant"]) in r["storage_path"]

    # Once: a later run neither stores it again nor finds it still owed.
    again = await _tick(w, ended + timedelta(minutes=8), out)
    assert again["stored"] == 0
    assert (await _sql("SELECT count(*) AS n FROM drone_reports WHERE session_id = :s", {"s": uuid.UUID(sid)}))[0]["n"] == 2
    assert out.sent == [], "nobody asked for this report by email"


# ─── C. Emails ───────────────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_immediate_report_goes_to_those_whose_scope_covers_the_flight(evidence):
    w = await _world(rules=[("intrusion", "low", None, "HIGH")])
    other_site = uuid.uuid4()
    await _sql("INSERT INTO sites (id, tenant_id, name) VALUES (:i,:t,'Elsewhere')", {"i": other_site, "t": w["tenant"]})
    at = _last(2, 17)
    sid, _ = await _flight_with_event(w, at)
    ended = at + timedelta(minutes=4)
    await _recipient(w, "everything@agency.test")
    await _recipient(w, "this-site@agency.test", site=w["site"])
    await _recipient(w, "this-mission@agency.test", mission=w["mission"])
    await _recipient(w, "other-site@agency.test", site=other_site)
    await _recipient(w, "daily@agency.test", "DAILY")
    await _recipient(w, "paused@agency.test", active=False)
    await _recipient(w, "added-later@agency.test", created=ended + timedelta(minutes=1))

    out = Outbox()
    got = await _tick(w, ended + timedelta(minutes=6), out)
    assert got["queued"] == 1 and got["sent"] == 1, got
    [mail] = out.sent
    assert mail["to"] == ["everything@agency.test", "this-mission@agency.test", "this-site@agency.test"]
    session_number = (await _sql("SELECT session_number FROM drone_patrol_sessions WHERE id = :s",
                                 {"s": uuid.UUID(sid)}))[0]["session_number"]
    assert "Night Watch" in mail["subject"] and session_number in mail["subject"] and "completed" in mail["subject"]
    assert mail["subtype"] == "pdf" and mail["payload"].startswith(b"%PDF")
    assert mail["filename"] == f"drone-patrol-{session_number}.pdf"
    assert "Events: 1 (1 suspicious). Incidents: 1." in mail["body"]
    [row] = await _queue(w)
    assert row["status"] == "SENT" and row["sent_at"] is not None and row["attempts"] == 1

    # Sent once: the next run has nothing to do.
    assert (await _tick(w, ended + timedelta(minutes=9), out))["sent"] == 0 and len(out.sent) == 1


@pytest.mark.asyncio
async def test_a_send_that_fails_is_retried_with_backoff_and_then_left_failed_with_its_reason(evidence):
    w = await _world()
    at = _last(2, 17)
    sid = await _flying(w, at, ended=True)
    await _recipient(w, "ops@agency.test")
    out = Outbox()
    out.down = True
    now = at + timedelta(minutes=6)

    first = await _tick(w, now, out)
    assert first["queued"] == 1 and first["failed"] == 1
    [row] = await _queue(w)
    assert row["status"] == "FAILED" and row["attempts"] == 1 and "refused" in row["last_error"]
    # The platform's own backoff, shared with Virtual Patrolling: five minutes
    # after the first failure, then fifteen, an hour, four hours.
    assert row["scheduled_at"] == now + timedelta(minutes=5)

    # Not before its time…
    assert (await _tick(w, now + timedelta(seconds=30), out))["failed"] == 0
    # …and then at growing intervals, until the attempts are spent.
    waits = [5, 15, 60, 240]
    for attempt, wait in enumerate(waits, start=2):
        now += timedelta(minutes=wait, seconds=1)
        assert (await _tick(w, now, out))["failed"] == 1
        assert (await _queue(w))[0]["attempts"] == attempt
    now += timedelta(hours=5)
    assert (await _tick(w, now, out)) == {"stored": 0, "queued": 0, "digests": 0, "sent": 0, "failed": 0}
    [row] = await _queue(w)
    assert row["status"] == "FAILED" and row["attempts"] == 5 and row["last_error"], "it must stay visible"

    # The mail server is fixed and someone asks for it again.
    out.down = False
    async with _client() as c:
        retry = await c.post(f"/api/v1/drone-report-deliveries/{row['id']}/retry", headers=w["h_admin"])
        log = await c.get("/api/v1/drone-report-deliveries", headers=w["h_admin"])
    assert retry.status_code == 200, retry.text
    assert log.json()["items"][0]["status"] == "PENDING" and log.json()["items"][0]["attempts_left"] == 5
    assert (await _tick(w, datetime.now(timezone.utc) + timedelta(seconds=5), out))["sent"] == 1
    assert out.sent[0]["to"] == ["ops@agency.test"] and (await _queue(w))[0]["status"] == "SENT"
    async with _client() as c:
        again = await c.post(f"/api/v1/drone-report-deliveries/{row['id']}/retry", headers=w["h_admin"])
    assert again.status_code == 409 and "sent" in again.json()["detail"]
    assert sid


@pytest.mark.asyncio
async def test_a_row_left_mid_send_by_a_dead_worker_is_taken_again(evidence):
    w = await _world()
    at = _last(2, 17)
    await _flying(w, at, ended=True)
    await _recipient(w, "ops@agency.test")
    out = Outbox()
    out.down = True
    now = at + timedelta(minutes=6)
    await _tick(w, now, out)
    await _sql("UPDATE drone_report_email_queue SET status = 'PROCESSING', claimed_at = :c WHERE tenant_id = :t",
               {"c": now, "t": w["tenant"]})
    out.down = False
    assert (await _tick(w, now + timedelta(minutes=10), out))["sent"] == 0, "too soon to call it dead"
    assert (await _tick(w, now + timedelta(minutes=31), out))["sent"] == 1


@pytest.mark.asyncio
async def test_a_summary_is_queued_once_per_scope_for_a_closed_period_with_flights_in_it(evidence):
    w = await _world()
    quiet_site = uuid.uuid4()
    await _sql("INSERT INTO sites (id, tenant_id, name) VALUES (:i,:t,'Quiet Yard')", {"i": quiet_site, "t": w["tenant"]})
    sid = await _flying(w, _last(2, 17), ended=True)
    now = datetime.now(timezone.utc)
    yesterday = now.astimezone(SGT).date() - timedelta(days=1)
    # The flight was created yesterday, at midday in the organisation's zone.
    await _sql("UPDATE drone_patrol_sessions SET created_at = :c, report_queued_at = now() WHERE id = :s",
               {"c": datetime.combine(yesterday, time(12, 0), tzinfo=SGT), "s": uuid.UUID(sid)})
    await _recipient(w, "director@agency.test", "DAILY")
    await _recipient(w, "deputy@agency.test", "DAILY")
    await _recipient(w, "site-lead@agency.test", "DAILY", site=w["site"])
    await _recipient(w, "quiet@agency.test", "DAILY", site=quiet_site)

    out = Outbox()
    got = await _tick(w, now, out)
    assert got["digests"] == 2 and got["sent"] == 2, got
    rows = {r["scope_key"]: r for r in await _queue(w)}
    assert set(rows) == {"tenant", f"site:{w['site']}"}, "the quiet site had nothing to report"
    assert rows["tenant"]["recipients"] == "deputy@agency.test,director@agency.test"
    assert rows["tenant"]["period_start"] == rows["tenant"]["period_end"] == yesterday
    assert rows["tenant"]["timezone"] == "Asia/Singapore"

    whole = next(m for m in out.sent if m["to"] == ["deputy@agency.test", "director@agency.test"])
    assert whole["subtype"] == XLSX and "daily summary: all sites" in whole["subject"]
    assert "Flights: 1 (1 completed, 0 did not complete)." in whole["body"]
    book = _sheets(whole["payload"])
    assert list(book) == ["Summary", "Flights", "Events"] and len(book["Flights"]) == 2
    assert dict((r[0], r[1]) for r in book["Summary"][1:])["Covers"] == "all sites"
    site_mail = next(m for m in out.sent if m["to"] == ["site-lead@agency.test"])
    assert "daily summary: Depot" in site_mail["subject"]

    # Once per period: later the same day nothing is queued or sent again.
    later = await _tick(w, now + timedelta(hours=2), out)
    assert later["digests"] == 0 and later["sent"] == 0 and len(await _queue(w)) == 2


@pytest.mark.asyncio
async def test_the_organisations_day_decides_which_summary_a_flight_is_in():
    w = await _world()
    day = datetime.now(SGT).date() - timedelta(days=3)
    late, after_midnight = uuid.uuid4(), uuid.uuid4()
    stmts = []
    for sid, local in ((late, datetime.combine(day, time(23, 30), tzinfo=SGT)),
                       (after_midnight, datetime.combine(day + timedelta(days=1), time(0, 30), tzinfo=SGT))):
        stmts.append(("INSERT INTO drone_patrol_sessions (id, tenant_id, session_number, site_id, mission_id, "
                      "    mission_name, status, started_at, ended_at, created_at) "
                      "VALUES (:i,:t,:n,:s,:m,'Night Watch','MISSED',:c,:c,:c)",
                      {"i": sid, "t": w["tenant"], "n": f"DP-{sid.hex[:10]}", "s": w["site"], "m": w["mission"],
                       "c": local}))
    await _run(stmts)
    async with AsyncSessionLocal() as db:
        await db.execute(text("SELECT set_config('app.current_tenant', :t, true)"), {"t": str(w["tenant"])})
        that_day = await reports.load_period(db, start=day, end=day)
        next_day = await reports.load_period(db, start=day + timedelta(days=1), end=day + timedelta(days=1))
    # 23:30 and 00:30 local are the same UTC date; they are not the same day.
    assert [f["id"] for f in that_day["flights"]] == [late]
    assert [f["id"] for f in next_day["flights"]] == [after_midnight]
    assert that_day["totals"]["by_status"] == {"MISSED": 1} and that_day["totals"]["did_not_complete"] == 1


# ─── D. Across tenants ───────────────────────────────────────────────────────

@pytest.mark.asyncio
async def test_the_report_job_runs_as_the_app_keeps_tenants_apart_and_serves_a_lapsed_licence(evidence):
    async with AsyncSessionLocal() as db:
        who = (await db.execute(text(
            "SELECT current_user, (SELECT rolbypassrls FROM pg_roles WHERE rolname = current_user)"))).first()
    assert who[1] is False, f"the report job connects as {who[0]!r}, which bypasses RLS"

    a, b = await _world(), await _world()
    at = _last(2, 17)
    sid_a, sid_b = await _flying(a, at, ended=True), await _flying(b, at, ended=True)
    await _recipient(a, "a@agency.test")
    await _recipient(b, "b@agency.test")
    # Tenant A's licence lapsed after its flight: the report is still owed.
    await _sql("UPDATE drone_module_licenses SET is_enabled = FALSE WHERE tenant_id = :t", {"t": a["tenant"]})

    out = Outbox()
    await runner.run_report_tick(AsyncSessionLocal, now=at + timedelta(minutes=6), deliver=out)
    numbers = {str(r["id"]): r["session_number"] for r in await _sql(
        "SELECT id, session_number FROM drone_patrol_sessions WHERE id = ANY(:ids)",
        {"ids": [uuid.UUID(sid_a), uuid.UUID(sid_b)]})}
    mine = {m["to"][0]: m for m in out.sent if m["to"][0] in ("a@agency.test", "b@agency.test")}
    assert set(mine) == {"a@agency.test", "b@agency.test"}, [m["to"] for m in out.sent]
    assert numbers[sid_a] in mine["a@agency.test"]["subject"] and numbers[sid_b] not in mine["a@agency.test"]["subject"]
    assert numbers[sid_b] in mine["b@agency.test"]["subject"]
    for w, sid in ((a, sid_a), (b, sid_b)):
        [row] = await _queue(w)
        assert row["session_id"] == uuid.UUID(sid) and row["status"] == "SENT"
        stored = await _sql("SELECT storage_path FROM drone_reports WHERE tenant_id = :t", {"t": w["tenant"]})
        assert len(stored) == 2 and all(str(w["tenant"]) in r["storage_path"] for r in stored)


# ─── E. The API ──────────────────────────────────────────────────────────────

async def _bare_session(w: dict, site: str = "site_a") -> uuid.UUID:
    sid = uuid.uuid4()
    await _sql("INSERT INTO drone_patrol_sessions (id, tenant_id, session_number, site_id, mission_name, drone_name, "
               "    status, started_at, launched_at, ended_at) "
               "VALUES (:i,:t,:n,:s,'Perimeter round','Drone One','COMPLETED', now() - interval '20 minutes', "
               "        now() - interval '20 minutes', now() - interval '8 minutes')",
               {"i": sid, "t": w["tenant"], "n": f"DP-{sid.hex[:10]}", "s": w[site]})
    return sid


@pytest.mark.asyncio
async def test_who_may_read_a_report_and_who_may_take_it_out(evidence):
    w, stranger = await _roles_world(), await _roles_world()
    here, there = await _bare_session(w), await _bare_session(w, "site_b")
    foreign = await _bare_session(stranger)
    base = f"/api/v1/drone-patrols/{here}/report"
    async with _client() as c:
        viewer_json = await c.get(base, headers=w["h"][VIEWER])
        viewer_pdf = await c.get(f"{base}/pdf", headers=w["h"][VIEWER])
        viewer_xlsx = await c.get(f"{base}/excel", headers=w["h"][VIEWER])
        admin_xlsx = await c.get(f"{base}/excel", headers=w["h"][ADMIN])
        guard_pdf = await c.get(f"{base}/pdf", headers=w["h"][GUARD])
        operator_xlsx = await c.get(f"{base}/excel", headers=w["h"][OPERATOR])
        out_of_scope = await c.get(f"/api/v1/drone-patrols/{there}/report/pdf", headers=w["h"][SUPERVISOR])
        in_scope = await c.get(f"{base}/pdf", headers=w["h"][SUPERVISOR])
        not_ours = await c.get(f"/api/v1/drone-patrols/{foreign}/report", headers=w["h"][ADMIN])
    assert viewer_json.status_code == 200 and viewer_json.json()["mission"]["mission_name"] == "Perimeter round"
    assert viewer_json.json()["stored"] == []
    assert viewer_pdf.status_code == 200 and viewer_pdf.content.startswith(b"%PDF")
    assert viewer_pdf.headers["content-type"] == "application/pdf"
    assert "attachment" in viewer_pdf.headers["content-disposition"] and ".pdf" in viewer_pdf.headers["content-disposition"]
    # Reading is not exporting.
    assert viewer_xlsx.status_code == 403 and operator_xlsx.status_code == 403
    assert admin_xlsx.status_code == 200 and admin_xlsx.content[:2] == b"PK"
    assert guard_pdf.status_code == 403
    assert out_of_scope.status_code == 404 and in_scope.status_code == 200
    assert not_ours.status_code == 404


@pytest.mark.asyncio
async def test_recipients_are_checked_scoped_and_can_always_be_stopped():
    w = await _roles_world()
    admin, sup = w["h"][ADMIN], w["h"][SUPERVISOR]
    url = "/api/v1/drone-report-recipients"
    async with _client() as c:
        made = await c.post(url, headers=admin, json={"email": " Ops@Agency.Test ", "frequency": "WEEKLY"})
        twice = await c.post(url, headers=admin, json={"email": "ops@agency.test", "frequency": "WEEKLY"})
        other_frequency = await c.post(url, headers=admin, json={"email": "ops@agency.test"})
        bad = await c.post(url, headers=admin, json={"email": "ops@agency.test\nBcc: x@evil.test"})
        both = await c.post(url, headers=admin, json={"email": "x@agency.test", "site_id": str(w["site_a"]),
                                                      "mission_id": str(uuid.uuid4())})
        viewer = await c.post(url, headers=w["h"][VIEWER], json={"email": "v@agency.test"})
        # A supervisor of site A only: their site yes, another site and "everything" no.
        own = await c.post(url, headers=sup, json={"email": "a@agency.test", "site_id": str(w["site_a"])})
        elsewhere = await c.post(url, headers=sup, json={"email": "b@agency.test", "site_id": str(w["site_b"])})
        everything = await c.post(url, headers=sup, json={"email": "c@agency.test"})
        seen_by_sup = await c.get(url, headers=sup)
        seen_by_admin = await c.get(url, headers=admin)
        hidden = await c.put(f"{url}/{made.json()['id']}", headers=sup, json={"is_active": False})
        paused = await c.put(f"{url}/{made.json()['id']}", headers=admin, json={"is_active": False})
    assert made.status_code == 201 and made.json()["email"] == "ops@agency.test", made.text
    assert made.json()["scope"] == "organisation" and made.json()["frequency"] == "WEEKLY"
    assert twice.status_code == 409 and other_frequency.status_code == 201
    assert bad.status_code == 422 and both.status_code == 422
    assert viewer.status_code == 403
    assert own.status_code == 201 and own.json()["scope"] == "site" and own.json()["site_name"] == "Factory A"
    assert elsewhere.status_code == 404 and everything.status_code == 422
    assert [r["email"] for r in seen_by_sup.json()] == ["a@agency.test"]
    assert len(seen_by_admin.json()) == 3
    assert hidden.status_code == 404 and paused.status_code == 200 and paused.json()["is_active"] is False

    # The licence lapses: nobody new can be added, but anyone can be removed.
    await _sql("UPDATE drone_module_licenses SET is_enabled = FALSE WHERE tenant_id = :t", {"t": w["tenant"]})
    async with _client() as c:
        add = await c.post(url, headers=admin, json={"email": "late@agency.test"})
        remove = await c.delete(f"{url}/{made.json()['id']}", headers=admin)
        left = await c.get(url, headers=admin)
    assert add.status_code == 403 and remove.status_code == 204
    assert made.json()["id"] not in [r["id"] for r in left.json()]
    audit = await _sql("SELECT action FROM audit_logs WHERE tenant_id = :t AND action LIKE 'drone.report_recipient.%' "
                       " ORDER BY created_at", {"t": w["tenant"]})
    assert {r["action"] for r in audit} == {"drone.report_recipient.create", "drone.report_recipient.update",
                                           "drone.report_recipient.delete"}


@pytest.mark.asyncio
async def test_a_period_summary_reads_and_exports_within_the_callers_sites():
    w = await _roles_world()
    await _bare_session(w)
    await _bare_session(w, "site_b")
    today = datetime.now(SGT).date()
    q = {"from": str(today - timedelta(days=1)), "to": str(today)}
    async with _client() as c:
        admin = await c.get("/api/v1/drone-reports/summary", headers=w["h"][ADMIN], params=q)
        sup = await c.get("/api/v1/drone-reports/summary", headers=w["h"][SUPERVISOR], params=q)
        sup_other = await c.get("/api/v1/drone-reports/summary", headers=w["h"][SUPERVISOR],
                                params={**q, "site_id": str(w["site_b"])})
        book = await c.get("/api/v1/drone-reports/summary/excel", headers=w["h"][ADMIN], params=q)
        viewer_book = await c.get("/api/v1/drone-reports/summary/excel", headers=w["h"][VIEWER], params=q)
        backwards = await c.get("/api/v1/drone-reports/summary", headers=w["h"][ADMIN],
                                params={"from": str(today), "to": str(today - timedelta(days=1))})
        too_long = await c.get("/api/v1/drone-reports/summary", headers=w["h"][ADMIN],
                               params={"from": str(today - timedelta(days=200)), "to": str(today)})
    assert admin.status_code == 200 and admin.json()["totals"]["flights"] == 2 and admin.json()["scope"] == "all sites"
    assert sup.json()["totals"]["flights"] == 1 and sup.json()["scope"] == "your sites"
    assert sup_other.status_code == 404
    assert book.status_code == 200 and len(_sheets(book.content)["Flights"]) == 3
    assert viewer_book.status_code == 403
    assert backwards.status_code == 422 and too_long.status_code == 422
